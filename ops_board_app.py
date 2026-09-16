"""Spreadsheet-style operations board for Bloomberg delivery work.

This is the practical replacement direction for the shipping Excel sheet:
new orders enter here first, rows are directly editable, and Excel becomes an
import/export format rather than the live work surface.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sqlite3
import threading
import webbrowser
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl

from config import EXCEL_PATH

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "ops_board.db"
EXPORT_DIR = Path.home() / "Downloads"

ORDER_RE = re.compile(r"\b(ZOR|ZRX|SDSK|ZINP|ZINT|ZRE|ORD)\s*[:#-]?\s*(\d{7,})\b", re.I)
OBD_RE = re.compile(r"\bOBD\s*[:#-]?\s*(\d{7,})\b", re.I)
SDSK_RE = re.compile(r"\bSDSK\s*[:#-]?\s*(\d{7,})\b", re.I)
ORD_RE = re.compile(r"\bORD\s*[:#-]?\s*(\d{7,})\b", re.I)
QTY_RE = re.compile(r"\bQty\s*[:#-]?\s*(\d+)\b", re.I)

COLUMNS = [
    "status", "date_text", "order_type", "order_no", "refs", "item", "material",
    "serial", "customer", "phone", "address", "client_memo",
]

LABELS = {
    "status": "Status",
    "date_text": "Date",
    "order_type": "Order Type",
    "order_no": "Order #",
    "refs": "OBD, ORD, SDSK",
    "item": "item",
    "material": "M/N",
    "serial": "S/N",
    "customer": "customer",
    "phone": "phone",
    "address": "ADDRESS",
    "client_memo": "Client Memo",
}

STATUS_OPTIONS = ["New", "Today", "Tomorrow", "Scheduled - Client", "Scheduled - Bloomberg", "Waiting", "Done", "Hold"]
COLORS = ["", "#e5e7eb", "#dbeafe", "#cffafe", "#dcfce7", "#fef3c7", "#fee2e2", "#f3e8ff", "#fce7f3"]


def now_text():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def weekday_ko(d: date):
    return "월화수목금토일"[d.weekday()]


def today_sheet_names(sheetnames):
    today = date.today()
    return [name for name in (f"{today.month}-{today.day}", f"{today.month}.{today.day}", f"{today.month}_{today.day}") if name in sheetnames]


def today_label():
    today = date.today()
    return f"{today.month}-{today.day}-{weekday_ko(today)}"


def date_label_from_header(value):
    text = text_value(value)
    match = re.search(r"(\d{4})년\s*(\d{1,2})월\s*(\d{1,2})일\s*([월화수목금토일])", text)
    if not match:
        return ""
    return f"{int(match.group(2))}-{int(match.group(3))}-{match.group(4)}"


def status_from_section(value):
    text = text_value(value)
    if not text:
        return ""
    lowered = text.lower()
    if "client" in lowered and "scheduled" in lowered:
        return "Scheduled - Client"
    if "bloomberg" in lowered and "scheduled" in lowered:
        return "Scheduled - Bloomberg"
    if "delayed" in lowered and "client" in lowered:
        return "Scheduled - Client"
    if "delayed" in lowered and "bloomberg" in lowered:
        return "Scheduled - Bloomberg"
    return ""


def date_label_from_any(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        d = value.date()
        return f"{d.month}-{d.day}-{weekday_ko(d)}"
    if isinstance(value, date):
        return f"{value.month}-{value.day}-{weekday_ko(value)}"
    text = text_value(value)
    label = date_label_from_header(text)
    if label:
        return label
    match = re.search(r"(\d{1,2})\s*/\s*(\d{1,2})", text)
    if match:
        month, day = int(match.group(1)), int(match.group(2))
        year = date.today().year
        try:
            d = date(year, month, day)
            return f"{d.month}-{d.day}-{weekday_ko(d)}"
        except ValueError:
            return ""
    return ""


def status_from_date_label(label):
    today = date.today()
    today_text = f"{today.month}-{today.day}-{weekday_ko(today)}"
    tomorrow = date.fromordinal(today.toordinal() + 1)
    tomorrow_text = f"{tomorrow.month}-{tomorrow.day}-{weekday_ko(tomorrow)}"
    if label == today_text:
        return "Today"
    if label == tomorrow_text:
        return "Tomorrow"
    return "Scheduled - Client"


def text_value(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def digits(value):
    return "".join(ch for ch in text_value(value) if ch.isdigit())


def parse_order_text(value):
    text = text_value(value)
    order_type = ""
    order_no = ""
    work_prefix = ""
    if "배송" in text:
        work_prefix = "배송"
    elif "회수" in text:
        work_prefix = "회수"
    match = ORDER_RE.search(text)
    if match:
        order_code, order_no = match.group(1).upper(), match.group(2)
        order_type = f"{work_prefix} {order_code}".strip()
    refs = []
    obd = OBD_RE.search(text)
    if obd:
        refs.append(f"OBD {obd.group(1)}")
    ord_match = ORD_RE.search(text)
    if ord_match:
        refs.append(f"ORD {ord_match.group(1)}")
    sdsk = SDSK_RE.search(text)
    if sdsk:
        refs.append(f"SDSK {sdsk.group(1)}")
    return order_type, order_no, "\n".join(refs)


def normalize_item(desc):
    desc = text_value(desc)
    qty = QTY_RE.search(desc)
    return desc


def ensure_db():
    with sqlite3.connect(DB_PATH) as con:
        con.execute("PRAGMA journal_mode=WAL")
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS board_rows (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sort_index REAL NOT NULL,
                completed INTEGER NOT NULL DEFAULT 0,
                status TEXT NOT NULL DEFAULT 'New',
                date_text TEXT DEFAULT '',
                order_type TEXT DEFAULT '',
                order_no TEXT DEFAULT '',
                refs TEXT DEFAULT '',
                item TEXT DEFAULT '',
                material TEXT DEFAULT '',
                serial TEXT DEFAULT '',
                customer TEXT DEFAULT '',
                phone TEXT DEFAULT '',
                address TEXT DEFAULT '',
                client_memo TEXT DEFAULT '',
                row_color TEXT DEFAULT '',
                cell_colors TEXT DEFAULT '{}',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        cols = {row[1] for row in con.execute("PRAGMA table_info(board_rows)")}
        if "cell_colors" not in cols:
            con.execute("ALTER TABLE board_rows ADD COLUMN cell_colors TEXT DEFAULT '{}'")
        if "completed" not in cols:
            con.execute("ALTER TABLE board_rows ADD COLUMN completed INTEGER NOT NULL DEFAULT 0")


def connect():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con


def row_dict(row):
    data = dict(row)
    try:
        data["cell_colors"] = json.loads(data.get("cell_colors") or "{}")
    except Exception:
        data["cell_colors"] = {}
    return data


def list_rows(query=""):
    ensure_db()
    where = ""
    params = []
    if query:
        where = "WHERE " + " OR ".join([f"{col} LIKE ?" for col in COLUMNS])
        params = [f"%{query}%"] * len(COLUMNS)
    with connect() as con:
        rows = con.execute(f"SELECT * FROM board_rows {where} ORDER BY sort_index ASC, id ASC", params).fetchall()
        return [row_dict(r) for r in rows]


def next_top_index(con):
    row = con.execute("SELECT MIN(sort_index) FROM board_rows").fetchone()
    current = row[0]
    return (current - 1) if current is not None else 0


def add_new_row():
    ensure_db()
    with connect() as con:
        cur = con.execute(
            """
            INSERT INTO board_rows(sort_index, status, date_text, created_at, updated_at)
            VALUES (?, 'New', ?, ?, ?)
            """,
            (next_top_index(con), today_label(), now_text(), now_text()),
        )
        return get_row(cur.lastrowid)


def get_row(row_id):
    with connect() as con:
        row = con.execute("SELECT * FROM board_rows WHERE id=?", (row_id,)).fetchone()
        return row_dict(row) if row else None


def update_cell(row_id, field, value):
    if field not in COLUMNS and field != "row_color":
        raise ValueError("invalid field")
    ensure_db()
    with connect() as con:
        con.execute(f"UPDATE board_rows SET {field}=?, updated_at=? WHERE id=?", (str(value or ""), now_text(), row_id))
        return get_row(row_id)


def toggle_completed(row_id, completed):
    ensure_db()
    with connect() as con:
        con.execute("UPDATE board_rows SET completed=?, updated_at=? WHERE id=?", (1 if completed else 0, now_text(), row_id))
        return get_row(row_id)


def update_cell_color(row_id, field, color):
    if field not in COLUMNS:
        raise ValueError("invalid field")
    ensure_db()
    with connect() as con:
        row = con.execute("SELECT cell_colors FROM board_rows WHERE id=?", (row_id,)).fetchone()
        if not row:
            raise ValueError("row not found")
        try:
            colors = json.loads(row["cell_colors"] or "{}")
        except Exception:
            colors = {}
        if color:
            colors[field] = color
        else:
            colors.pop(field, None)
        con.execute("UPDATE board_rows SET cell_colors=?, updated_at=? WHERE id=?", (json.dumps(colors, ensure_ascii=False), now_text(), row_id))
        return get_row(row_id)


def reorder_rows(ids):
    ensure_db()
    with connect() as con:
        for idx, row_id in enumerate(ids):
            con.execute("UPDATE board_rows SET sort_index=?, updated_at=? WHERE id=?", (idx, now_text(), int(row_id)))
    return {"ok": True}


def delete_row(row_id):
    with connect() as con:
        con.execute("DELETE FROM board_rows WHERE id=?", (row_id,))
    return {"ok": True}


def clear_rows():
    with connect() as con:
        con.execute("DELETE FROM board_rows")
    return {"ok": True}


def import_today_from_excel(replace=False, sheet_name=None):
    ensure_db()
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    sheets = [sheet_name] if sheet_name else today_sheet_names(wb.sheetnames)
    sheets = [name for name in sheets if name in wb.sheetnames]
    if not sheets:
        wb.close()
        raise RuntimeError("target sheet not found")
    inserted = 0
    try:
        with connect() as con:
            if replace:
                con.execute("DELETE FROM board_rows")
            max_sort = con.execute("SELECT COALESCE(MAX(sort_index), -1) FROM board_rows").fetchone()[0]
            sort_index = max_sort + 1
            for sheet_name in sheets:
                ws = wb[sheet_name]
                current_date = ""
                current_status = "Today"
                carry = {"order_type": "", "order_no": "", "refs": "", "customer": "", "phone": "", "address": "", "client_memo": ""}
                for row_num in range(1, ws.max_row + 1):
                    date_header = date_label_from_header(ws.cell(row_num, 7).value)
                    if date_header:
                        current_date = date_header
                        current_status = status_from_date_label(current_date)
                        carry = {"order_type": "", "order_no": "", "refs": "", "customer": "", "phone": "", "address": "", "client_memo": ""}
                        continue
                    section_status = status_from_section(ws.cell(row_num, 1).value)
                    if section_status:
                        current_status = section_status
                        carry = {"order_type": "", "order_no": "", "refs": "", "customer": "", "phone": "", "address": "", "client_memo": ""}
                        continue

                    a = text_value(ws.cell(row_num, 1).value)
                    item = normalize_item(ws.cell(row_num, 2).value)
                    material = digits(ws.cell(row_num, 3).value)
                    serial = text_value(ws.cell(row_num, 4).value)
                    customer = text_value(ws.cell(row_num, 5).value)
                    phone = text_value(ws.cell(row_num, 6).value)
                    address = text_value(ws.cell(row_num, 7).value)
                    memo = text_value(ws.cell(row_num, 8).value)
                    status_note = text_value(ws.cell(row_num, 9).value)

                    if item.lower() == "item" or text_value(ws.cell(row_num, 3).value).upper() in {"M/N", "MN", "MATERIAL"}:
                        continue
                    if a:
                        order_type, order_no, refs = parse_order_text(a)
                        if order_no:
                            carry["order_type"] = order_type
                            carry["order_no"] = order_no
                            carry["refs"] = refs
                    for key, value in (("customer", customer), ("phone", phone), ("address", address), ("client_memo", memo)):
                        if value:
                            carry[key] = value
                    if not any([item, material, serial, carry.get("order_no"), customer, phone, address, memo, status_note]):
                        continue
                    if not any([item, material, serial]):
                        continue
                    note_date = date_label_from_any(ws.cell(row_num, 9).value)
                    if note_date:
                        row_date = note_date
                        row_status = current_status or status_from_date_label(note_date)
                    else:
                        row_date = current_date or sheet_name
                        row_status = status_note or current_status or status_from_date_label(row_date)
                    con.execute(
                        """
                        INSERT INTO board_rows(sort_index, completed, status, date_text, order_type, order_no, refs, item,
                            material, serial, customer, phone, address, client_memo, row_color, cell_colors, created_at, updated_at)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            sort_index, 0, row_status, row_date, carry.get("order_type", ""), carry.get("order_no", ""),
                            carry.get("refs", ""), item, material, serial, carry.get("customer", ""), carry.get("phone", ""),
                            carry.get("address", ""), carry.get("client_memo", ""), "#e5e7eb", "{}", now_text(), now_text(),
                        ),
                    )
                    sort_index += 1
                    inserted += 1
    finally:
        wb.close()
    return {"inserted": inserted, "sheets": sheets}


def export_xlsx():
    ensure_db()
    rows = list_rows()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Board"
    headers = [LABELS[c] for c in COLUMNS]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(c, "") for c in COLUMNS])
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 16
    ws.column_dimensions["F"].width = 34
    ws.column_dimensions["K"].width = 36
    ws.column_dimensions["L"].width = 24
    filename = EXPORT_DIR / f"ops_board_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    return {"path": str(filename), "rows": len(rows)}


HTML = r'''
<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Operations Board</title>
<style>
:root{font-family:"Malgun Gothic",Arial,sans-serif;color:#111827;background:#f3f4f6}body{margin:0}.top{height:48px;background:#20242c;color:#fff;display:flex;align-items:center;justify-content:space-between;padding:0 14px}.top h1{font-size:17px;margin:0}.toolbar{display:flex;gap:7px;align-items:center;padding:10px 12px;background:#fff;border-bottom:1px solid #cfd5dd;position:sticky;top:0;z-index:5}.btn,input{height:32px;border:1px solid #aab2bd;border-radius:3px;background:#fff;padding:0 9px;font:13px "Malgun Gothic",Arial}.btn{cursor:pointer}.primary{background:#0f62fe;color:white;border-color:#0f62fe}.danger{background:#fee2e2}.palette{display:flex;gap:4px;align-items:center;border-left:1px solid #d1d5db;padding-left:8px}.swatch{width:22px;height:22px;border:1px solid #94a3b8;border-radius:2px;cursor:pointer}.board-wrap{padding:12px;overflow:auto;height:calc(100vh - 101px)}table{border-collapse:separate;border-spacing:0;min-width:1270px;width:100%;table-layout:fixed}th,td{border-right:1px solid #fff;border-bottom:1px solid #fff;background:#e5e7eb;font-size:12px;vertical-align:top}th{height:24px;background:#dfe3e8;text-align:center;font-weight:500;position:sticky;top:0;z-index:2}td{height:54px;padding:0}.cell{min-height:54px;padding:6px;white-space:pre-wrap;word-break:break-word;outline:none}.cell:focus{box-shadow:inset 0 0 0 2px #0f62fe;background:#eef6ff}.selected{box-shadow:inset 0 0 0 2px #111827}.range-selected{box-shadow:inset 0 0 0 2px #2563eb;background:#dbeafe!important}.drag{width:32px;text-align:center;background:#d1d5db;cursor:grab;color:#475569;font-weight:700}.donecol{width:34px;text-align:center;background:#d1d5db}.donebox{width:16px;height:16px;margin-top:18px}.dragging{opacity:.45}.drop-target td{box-shadow:inset 0 2px 0 #0f62fe}.status{width:86px;text-align:center}.date{width:86px}.otype{width:80px}.order{width:58px}.refs{width:120px}.item{width:200px}.mn{width:76px}.sn{width:70px}.customer{width:78px}.phone{width:62px}.address{width:202px}.memo{width:94px}.hint{color:#64748b;font-size:12px;margin-left:auto}.toast{position:fixed;right:14px;bottom:14px;background:#111827;color:#fff;padding:9px 12px;border-radius:4px;display:none}.newrow td{background:#eef2ff}.completed-row td{background:#bfdbfe!important}.delete{width:34px;text-align:center}.mini{height:24px;padding:0 7px}</style>
</head>
<body>
<div class="top"><h1>Operations Board</h1><div>체크하면 완료, 셀 드래그 선택 후 Ctrl+C로 Excel 복사</div></div>
<div class="toolbar">
<button class="btn primary" onclick="addRow()">새 오더 행</button>
<button class="btn" onclick="importToday(false)">오늘 Excel 불러오기</button>
<button class="btn danger" onclick="importToday(true)">비우고 다시 불러오기</button>
<button class="btn" onclick="exportExcel()">Excel로 내보내기</button>
<input id="q" placeholder="검색" oninput="loadRows()">
<div class="palette" id="palette"></div>
<button class="btn" onclick="copySelection()">선택 범위 복사</button>
<span class="hint">마우스로 여러 칸 드래그 선택, Ctrl+C 복사, Ctrl+Enter 줄바꿈, Tab 이동</span>
</div>
<div class="board-wrap"><table id="board"><thead><tr id="head"></tr></thead><tbody id="body"></tbody></table></div>
<div class="toast" id="toast"></div>
<script>
const cols=["status","date_text","order_type","order_no","refs","item","material","serial","customer","phone","address","client_memo"];
const labels={status:"Status",date_text:"Date",order_type:"Order Type",order_no:"Order #",refs:"OBD, ORD, SDSK",item:"item",material:"M/N",serial:"S/N",customer:"customer",phone:"phone",address:"ADDRESS",client_memo:"Client Memo"};
const classes={status:"status",date_text:"date",order_type:"otype",order_no:"order",refs:"refs",item:"item",material:"mn",serial:"sn",customer:"customer",phone:"phone",address:"address",client_memo:"memo"};
const colors=["","#e5e7eb","#dbeafe","#cffafe","#dcfce7","#fef3c7","#fee2e2","#f3e8ff","#fce7f3","#ffffff","#fde68a","#a7f3d0","#bfdbfe","#fbcfe8"];
const mergeCols=new Set(["customer","phone","address","client_memo"]);
let selected=null;let rows=[];let draggedId=null;let isSelecting=false;let anchor=null;let range=[];
function toast(t){const el=document.getElementById('toast');el.textContent=t;el.style.display='block';setTimeout(()=>el.style.display='none',1700)}
function esc(s){return String(s??'').replace(/[&<>]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;'}[c]))}
async function api(path,opts={}){const r=await fetch(path,{headers:{'Content-Type':'application/json'},...opts});const d=await r.json();if(!r.ok)throw new Error(d.error||'request failed');return d}
function renderPalette(){document.getElementById('palette').innerHTML=colors.map(cl=>`<span class="swatch" title="${cl||'clear'}" style="background:${cl||'linear-gradient(135deg,#fff 45%,#ef4444 47%,#ef4444 53%,#fff 55%)'}" onclick="paintSelected('${cl}')"></span>`).join('')}
function renderHead(){document.getElementById('head').innerHTML='<th class="drag">↕</th><th class="donecol">✓</th>'+cols.map(c=>`<th class="${classes[c]}">${labels[c]}</th>`).join('')+'<th class="delete">Del</th>'}
function carriedValue(idx,c){for(let i=idx;i>=0;i--){const v=String(rows[i]?.[c]??'').trim();if(v)return v}return ''}
function mergePlan(){const plan={};for(const c of mergeCols){plan[c]={skip:{},span:{},value:{}};let i=0;while(i<rows.length){const value=carriedValue(i,c);if(!value){i++;continue}let j=i+1;while(j<rows.length&&carriedValue(j,c)===value)j++;const span=j-i;plan[c].span[i]=span;plan[c].value[i]=value;for(let k=i+1;k<j;k++)plan[c].skip[k]=true;i=j}}return plan}
function cell(row,rowIdx,c,colIdx,plan){if(plan&&plan[c]?.skip[rowIdx])return '';const bg=(row.cell_colors&&row.cell_colors[c])||'';const span=plan&&plan[c]?.span[rowIdx]>1?` rowspan="${plan[c].span[rowIdx]}"`:'';const value=plan&&plan[c]?.value[rowIdx]?plan[c].value[rowIdx]:(row[c]||'');return `<td class="${classes[c]}"${span} style="background:${bg||'#e5e7eb'}"><div class="cell" contenteditable="true" data-id="${row.id}" data-row="${rowIdx}" data-col="${colIdx}" data-field="${c}" onclick="selectCell(this)" onmousedown="startSelect(event,this)" onmouseenter="extendSelect(event,this)" onblur="saveCell(this)" onkeydown="cellKey(event,this)">${esc(value)}</div></td>`}
function render(){const body=document.getElementById('body');const plan=mergePlan();body.innerHTML=rows.map((r,ri)=>`<tr draggable="true" data-id="${r.id}" class="${r.status==='New'?'newrow':''} ${r.completed?'completed-row':''}" ondragstart="dragStart(event,${r.id})" ondragover="dragOver(event,this)" ondragleave="this.classList.remove('drop-target')" ondrop="dropRow(event,this)" ondragend="dragEnd()"><td class="drag">::</td><td class="donecol"><input class="donebox" type="checkbox" ${r.completed?'checked':''} onchange="toggleDone(${r.id},this.checked)"></td>${cols.map((c,ci)=>cell(r,ri,c,ci,plan)).join('')}<td class="delete"><button class="btn mini" onclick="delRow(${r.id})">X</button></td></tr>`).join('')}
async function loadRows(){const q=document.getElementById('q').value;const d=await api('/api/rows?q='+encodeURIComponent(q));rows=d.rows;render()}
async function addRow(){const d=await api('/api/rows',{method:'POST'});await loadRows();toast('새 행 추가');setTimeout(()=>{const el=document.querySelector(`[data-id="${d.id}"][data-field="order_no"]`);if(el){el.focus();selectCell(el)}},50)}
async function importToday(replace){if(replace&&!confirm('현재 보드 행을 비우고 Excel을 다시 불러올까요?'))return;const d=await api('/api/import_today',{method:'POST',body:JSON.stringify({replace})});await loadRows();toast(`불러옴: ${d.inserted} rows`)}
async function exportExcel(){const d=await api('/api/export',{method:'POST'});toast('Export: '+d.path);alert(d.path)}
async function saveCell(el){await api(`/api/rows/${el.dataset.id}/cell`,{method:'POST',body:JSON.stringify({field:el.dataset.field,value:el.innerText})})}
function clearRange(){document.querySelectorAll('.range-selected').forEach(x=>x.classList.remove('range-selected'));range=[]}
function selectCell(el){document.querySelectorAll('.selected').forEach(x=>x.classList.remove('selected'));el.classList.add('selected');selected=el;if(!isSelecting){clearRange();range=[el];el.classList.add('range-selected')}}
function startSelect(e,el){if(e.button!==0)return;isSelecting=true;anchor=el;clearRange();selectCell(el)}
function extendSelect(e,el){if(!isSelecting||!anchor)return;markRange(anchor,el)}
function markRange(a,b){clearRange();const r1=Math.min(Number(a.dataset.row),Number(b.dataset.row));const r2=Math.max(Number(a.dataset.row),Number(b.dataset.row));const c1=Math.min(Number(a.dataset.col),Number(b.dataset.col));const c2=Math.max(Number(a.dataset.col),Number(b.dataset.col));for(const cell of document.querySelectorAll('.cell')){const r=Number(cell.dataset.row),c=Number(cell.dataset.col);if(r>=r1&&r<=r2&&c>=c1&&c<=c2){cell.classList.add('range-selected');range.push(cell)}}}
document.addEventListener('mouseup',()=>{isSelecting=false;anchor=null});
async function paintSelected(color){const targets=range.length?range:(selected?[selected]:[]);if(!targets.length){toast('먼저 칸을 선택');return}for(const el of targets){await api(`/api/rows/${el.dataset.id}/color`,{method:'POST',body:JSON.stringify({field:el.dataset.field,color})});el.parentElement.style.background=color||'#e5e7eb'}toast('색 적용')}
async function toggleDone(id,checked){await api(`/api/rows/${id}/completed`,{method:'POST',body:JSON.stringify({completed:checked})});await loadRows()}
async function delRow(id){if(!confirm('이 행 삭제?'))return;await api(`/api/rows/${id}`,{method:'DELETE'});await loadRows()}
function copySelection(){const targets=range.length?range:(selected?[selected]:[]);if(!targets.length)return;const byRow={};for(const el of targets){const r=Number(el.dataset.row),c=Number(el.dataset.col);if(!byRow[r])byRow[r]={};byRow[r][c]=el.innerText}const rowNums=Object.keys(byRow).map(Number).sort((a,b)=>a-b);const colNums=targets.map(el=>Number(el.dataset.col));const minC=Math.min(...colNums),maxC=Math.max(...colNums);const text=rowNums.map(r=>{let vals=[];for(let c=minC;c<=maxC;c++){vals.push(byRow[r][c]??'')}return vals.join('\t')}).join('\n');copyText(text)}
async function copyText(t){await navigator.clipboard.writeText(t);toast('복사됨')}
function cellKey(e,el){if(e.key==='Enter'&&e.ctrlKey){document.execCommand('insertLineBreak');e.preventDefault();return}if(e.key==='Tab'){e.preventDefault();saveCell(el);const cells=[...document.querySelectorAll('.cell')];const i=cells.indexOf(el);const next=cells[i+(e.shiftKey?-1:1)];if(next){next.focus();selectCell(next)}}}
function dragStart(e,id){if(e.target.classList.contains('cell')){e.preventDefault();return}draggedId=id;e.currentTarget.classList.add('dragging');e.dataTransfer.effectAllowed='move'}
function dragOver(e,tr){e.preventDefault();if(String(draggedId)!==tr.dataset.id)tr.classList.add('drop-target')}
async function dropRow(e,tr){e.preventDefault();document.querySelectorAll('.drop-target').forEach(x=>x.classList.remove('drop-target'));const targetId=Number(tr.dataset.id);if(!draggedId||draggedId===targetId)return;const from=rows.findIndex(r=>r.id===draggedId);const to=rows.findIndex(r=>r.id===targetId);const [moved]=rows.splice(from,1);rows.splice(to,0,moved);render();await api('/api/reorder',{method:'POST',body:JSON.stringify({ids:rows.map(r=>r.id)})});toast('행 순서 저장')}
function dragEnd(){document.querySelectorAll('.dragging,.drop-target').forEach(x=>x.classList.remove('dragging','drop-target'));draggedId=null}
document.addEventListener('keydown',e=>{if(e.ctrlKey&&e.key.toLowerCase()==='c'&&(range.length||selected)){e.preventDefault();copySelection()}});
renderPalette();renderHead();loadRows();
</script>
</body></html>
'''



class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        return

    def send_json(self, data, status=200):
        payload = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def read_json(self):
        length = int(self.headers.get("Content-Length", "0") or 0)
        return json.loads(self.rfile.read(length).decode("utf-8")) if length else {}

    def do_GET(self):
        try:
            parsed = urlparse(self.path)
            if parsed.path == "/":
                payload = HTML.encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)
                return
            if parsed.path == "/api/rows":
                q = parse_qs(parsed.query).get("q", [""])[0]
                self.send_json({"rows": list_rows(q)})
                return
            self.send_json({"error": "not found"}, 404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 500)

    def do_POST(self):
        try:
            parsed = urlparse(self.path)
            if parsed.path == "/api/rows":
                self.send_json(add_new_row())
                return
            if parsed.path == "/api/import_today":
                body = self.read_json()
                self.send_json(import_today_from_excel(replace=bool(body.get("replace"))))
                return
            if parsed.path == "/api/export":
                self.send_json(export_xlsx())
                return
            if parsed.path == "/api/reorder":
                body = self.read_json()
                self.send_json(reorder_rows(body.get("ids") or []))
                return
            m = re.match(r"^/api/rows/(\d+)/completed$", parsed.path)
            if m:
                body = self.read_json()
                self.send_json(toggle_completed(int(m.group(1)), bool(body.get("completed"))))
                return
            m = re.match(r"^/api/rows/(\d+)/color$", parsed.path)
            if m:
                body = self.read_json()
                self.send_json(update_cell_color(int(m.group(1)), body.get("field"), body.get("color")))
                return
            m = re.match(r"^/api/rows/(\d+)/cell$", parsed.path)
            if m:
                body = self.read_json()
                self.send_json(update_cell(int(m.group(1)), body.get("field"), body.get("value")))
                return
            self.send_json({"error": "not found"}, 404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 500)

    def do_DELETE(self):
        try:
            m = re.match(r"^/api/rows/(\d+)$", urlparse(self.path).path)
            if m:
                self.send_json(delete_row(int(m.group(1))))
                return
            self.send_json({"error": "not found"}, 404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 500)


def serve(host, port, open_browser=True):
    ensure_db()
    server = ThreadingHTTPServer((host, port), Handler)
    url = f"http://{host}:{port}/"
    print(f"Operations Board running: {url}", flush=True)
    if open_browser:
        threading.Timer(0.7, lambda: webbrowser.open(url)).start()
    server.serve_forever()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8770)
    parser.add_argument("--no-open", action="store_true")
    parser.add_argument("--import-today", action="store_true")
    parser.add_argument("--sheet")
    parser.add_argument("--replace", action="store_true")
    args = parser.parse_args()
    ensure_db()
    if args.import_today:
        print(json.dumps(import_today_from_excel(replace=args.replace, sheet_name=args.sheet), ensure_ascii=False, indent=2))
        return
    serve(args.host, args.port, open_browser=not args.no_open)


if __name__ == "__main__":
    main()
