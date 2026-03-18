"""
Excel 테스트: 실제 배송장.xlsx에 더미 데이터를 삽입해봅니다.
결과를 보고 삽입 위치와 포맷이 올바른지 확인하세요.

실행: python test_excel.py
"""

import sys
sys.path.insert(0, '.')

from excel_handler import (
    write_orders_to_excel, find_today_section_end, get_today_sheet_name
)
import openpyxl
from config import EXCEL_PATH


def inspect_excel():
    """Excel 파일 구조 파악."""
    wb = openpyxl.load_workbook(EXCEL_PATH)
    print(f"[파일] {EXCEL_PATH}")
    print(f"[시트 목록] {wb.sheetnames}")
    print(f"[활성 시트] {wb.active.title}")
    print()

    ws = wb.active
    print(f"[활성 시트 '{ws.title}'] 상위 25행:")
    for row in ws.iter_rows(min_row=1, max_row=25):
        for cell in row:
            if cell.value:
                print(f"  행{cell.row} 열{cell.column} ({cell.column_letter}): '{str(cell.value)[:80]}'")
    print()

    today_sheet = get_today_sheet_name()
    print(f"[오늘 시트명 추정] '{today_sheet}'")

    if today_sheet in wb.sheetnames:
        ws_today = wb[today_sheet]
    else:
        ws_today = wb.active
        print(f"  → 오늘 시트 없음, 활성 시트 사용: '{ws_today.title}'")

    end_row = find_today_section_end(ws_today)
    print(f"[삽입 위치] 오늘 섹션 마지막 데이터 행: {end_row} → 새 데이터는 {end_row+1}행부터 입력됨")


def test_insert():
    """더미 데이터 삽입 테스트."""
    print("\n[더미 데이터 삽입 테스트]")
    dummy_rows = [
        {
            'order_prefix': '배송',
            'order_type': 'ZOR',
            'order_num': '7999999',
            'extra_orders': ['SDSK 001234567'],
            'material': '10045196',
            'description': 'TEST ITEM - AUTO INSERT',
            'qty': 1,
            'customer': 'TEST USER',
            'phone': '010-1234-5678',
            'company': 'TEST COMPANY LTD',
            'street': '123 TEST STREET',
            'street2': 'GANGNAM-GU',
            'memo': '[테스트] 자동 삽입 확인용',
            'is_first_item': True,
        },
        {
            'order_prefix': '배송',
            'order_type': 'ZOR',
            'order_num': '7999999',
            'extra_orders': ['SDSK 001234567'],
            'material': '10045196',
            'description': 'TEST ITEM - AUTO INSERT',
            'qty': 1,
            'customer': '',
            'phone': '',
            'company': '',
            'street': '',
            'street2': '',
            'memo': '',
            'is_first_item': False,
        },
    ]

    confirm = input("실제 배송장.xlsx에 테스트 데이터를 삽입합니다. 계속할까요? (y/n): ")
    if confirm.lower() != 'y':
        print("취소됨")
        return

    write_orders_to_excel(dummy_rows)
    print("삽입 완료. 배송장.xlsx를 열어서 결과를 확인하세요.")
    print("확인 후 삽입된 TEST 행들을 수동으로 삭제해 주세요.")


if __name__ == "__main__":
    inspect_excel()
    test_insert()
