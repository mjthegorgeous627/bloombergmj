"""Excel <-> workbench sync: importing SAP/portal scan rows into the DB,
and exporting/mirroring finalized workbench state back out to Excel.

Split out of workbench_app.py on 2026-09-16 - see db.py's module docstring
and BLOOMBERG_HANDOFF.md. Pure code motion.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import excel_handler
from config import EXCEL_PATH, FINALIZE_EXPORT_PATH
from db import (
    BASE_DIR,
    _pop_prior_item_state,
    add_event,
    connect,
    digits,
    ensure_db,
    extract_qty,
    logger,
    now_text,
    parse_item_type,
    parse_order_text,
    row_date_header,
    text_value,
    today_sheet_names,
)

def _scan_sheet_groups(ws, sheet_name, row_start=1, row_end=None):
    """Parse one Excel sheet (or an explicit [row_start, row_end] slice of
    it) into order/item groups, same shape load_today_from_excel() used to
    build inline. Shared by load_today_from_excel() (whole "today" sheet(s),
    the old automatic path) and import_excel_range() (an explicit sheet +
    row range, the manual emergency-only tool - see both callers' own
    docstrings for why this is now manual-only, not part of the auto-sync
    loop). Date-header tracking (row_date_header()) still scans from row 1
    even when row_start is later, so a scoped range still gets the right
    date if its own rows don't repeat the header - callers that want a
    literal "just these rows, whatever date they claim" should pass
    row_start=1 if that matters."""
    row_end = row_end if row_end is not None else ws.max_row
    current = None
    current_date_iso = date.today().isoformat()  # until the first date-block header is seen
    groups = []
    for row in range(1, row_end + 1):
        header_date = row_date_header(ws, row)
        if header_date:
            current_date_iso = header_date
            current = None  # next order line starts a fresh group under this date
            continue

        # Order/customer identity tracking always runs from row 1, regardless
        # of row_start - a scoped range (import_excel_range()) commonly asks
        # for just the item rows of a block (e.g. rows 42-43) whose own
        # order-number/customer row sits a row or two earlier (row 41) and
        # wouldn't otherwise repeat there. Only the actual item line below is
        # gated by row_start, so "rows 42-43" correctly still attaches to the
        # right order/customer instead of silently importing nothing (a real
        # bug hit in production - see project_sap_watchdog_fix memory).
        order_text = text_value(ws.cell(row, 1).value)
        desc = text_value(ws.cell(row, 2).value)
        material = digits(ws.cell(row, 3).value)
        serial = text_value(ws.cell(row, 4).value)
        customer = text_value(ws.cell(row, 5).value)
        phone = text_value(ws.cell(row, 6).value)
        address = text_value(ws.cell(row, 7).value)
        memo = text_value(ws.cell(row, 8).value)
        item_type = parse_item_type(order_text)

        if order_text:
            order_type, order_no, delivery_no = parse_order_text(order_text)
            if order_no:
                current = {
                    "order_no": order_no,
                    "delivery_no": delivery_no,
                    "order_type": order_type,
                    "source_sheet": sheet_name,
                    "source_date": current_date_iso,
                    "source_row_start": row,
                    "source_row_end": row,
                    "customer": customer,
                    "phone": phone,
                    "address": address,
                    "memo": memo,
                    "items": [],
                }
                groups.append(current)
        if not current:
            continue
        if not any([desc, material, serial, customer, phone, address, memo]):
            continue
        if order_text:
            _, _, row_delivery = parse_order_text(order_text)
            if row_delivery:
                current["delivery_no"] = row_delivery
        current["source_row_end"] = row
        for key, value in [("customer", customer), ("phone", phone), ("address", address), ("memo", memo)]:
            if value and not current.get(key):
                current[key] = value
        is_header_item = (desc.strip().lower() == "item" and serial.strip().upper() in {"S/N", "SN", "SERIAL"})
        is_header_material = text_value(ws.cell(row, 3).value).strip().upper() in {"M/N", "MN", "MATERIAL"}
        if is_header_item or is_header_material:
            continue
        if row < row_start:
            continue  # identity/context above is still tracked; only the item itself is skipped
        if material or desc or serial:
            qty = extract_qty(desc)
            # Default is serial. Material-only is a rare exception; qty>1 with no S/N is the practical signal.
            mode = "material_only" if (not serial and material and qty > 1) else "serial"
            current["items"].append(
                {
                    "description": desc,
                    "material": material,
                    "serial": serial,
                    "qty": qty,
                    "mode": mode,
                    "item_type": item_type,
                }
            )
    return groups


def _ingest_groups(con, groups, additive=False):
    """Upsert already-parsed order/item groups (from _scan_sheet_groups())
    into workbench.db. Default (additive=False, load_today_from_excel()'s
    whole-sheet reimport) deletes and reinserts every matched order's items
    wholesale, which is why manual_item_edits/item_date/board_status/
    deleted_at/cell_highlights all need their own explicit carry-over logic
    below - see each's own comment. That wholesale replace is only safe
    when the scan actually covers an order's *complete* row range (true for
    a whole-sheet scan), and this used to be reused as-is by
    import_excel_range() too.

    Real bug found 2026-08-07: a scoped import_excel_range() scan (e.g.
    "just rows 27-28") only ever sees a *partial slice* of a sheet - if
    those rows happen to belong to an order that already exists in
    workbench.db (a real case: ZOR 7825897, order id 57, carried forward
    from its original 8/4 collection into 8/7's sheet - the user asked for
    just 2 of its rows), the old code still ran the same wholesale
    DELETE-all-items-then-reinsert-only-what-this-scan-saw path, silently
    destroying every item of that order NOT in the requested 2-row window -
    with no soft-delete, no undo-log entry (unlike every other
    data-mutating action in this file), so it wasn't even Ctrl+Z-recoverable
    afterward. additive=True (used by import_excel_range() - see its own
    docstring) fixes this: for an existing order, never deletes anything -
    only adds items from the scan that don't already match an existing
    (material, serial) item, same "never touch a field once written"
    guarantee import_sap_rows() already gives the routine SAP path.
    Returns the number of order-groups upserted."""
    imported = 0
    merged_groups = {}
    for group in groups:
        if not group["items"]:
            continue
        key = (group.get("order_no") or "", group.get("delivery_no") or "")
        if key not in merged_groups:
            merged_groups[key] = group
            continue
        target = merged_groups[key]
        target["items"].extend(group["items"])
        target["source_row_end"] = max(target.get("source_row_end") or 0, group.get("source_row_end") or 0)
        for field in ("customer", "phone", "address", "memo", "order_type", "source_date"):
            if group.get(field) and not target.get(field):
                target[field] = group[field]

    for group in merged_groups.values():
        group_date = group.get("source_date") or date.today().isoformat()
        existing = con.execute(
            "SELECT id, status FROM orders WHERE order_no=? AND COALESCE(delivery_no,'')=COALESCE(?, '')",
            (group["order_no"], group.get("delivery_no") or ""),
        ).fetchone()
        prior_by_key = {}
        prior_by_line = {}
        manual_edits_by_item = {}
        if existing and additive:
            # Scoped/partial scan (import_excel_range()) - this group may
            # only be a slice of the order's real full item list, so unlike
            # the branch below, never delete anything. Only bump
            # updated_at; leave order_type/source_sheet/source_date etc.
            # alone too (a partial scan's own row numbers/date don't
            # describe the order as a whole and shouldn't overwrite
            # whatever it already has).
            order_id = existing["id"]
            # Counter, not a set: several existing items can legitimately
            # share one (material, serial="") key (e.g. serial-less
            # material-only lines) - each incoming item should only match
            # off ONE such existing item, same one-to-one consumption
            # _pop_prior_item_state() already does for the non-additive
            # path below. Otherwise N>1 genuinely-new items sharing a key
            # with each other (not with anything pre-existing) would
            # wrongly dedupe against each other and under-recover.
            # Real bug found 2026-08-07 (order 67072686, sheet 8-10 rows
            # 19-20): this used to count ALL items regardless of deleted_at,
            # so an item a user had soft-deleted (still sitting in the 7-day
            # trash, e.g. from an earlier "체크된 행 삭제") permanently blocked
            # ever re-adding it via a scoped range import - the whole point
            # of the retry was to bring that exact item back, and it silently
            # matched against the deleted row and got skipped again. A
            # soft-deleted item is hidden, not "still on the board" - only
            # live (non-deleted) items should count as already-present.
            existing_counts = Counter()
            for prow in con.execute(
                "SELECT material, serial FROM order_items WHERE order_id=? AND deleted_at IS NULL", (order_id,)
            ).fetchall():
                existing_counts[(prow["material"] or "", prow["serial"] or "")] += 1
            next_line = con.execute(
                "SELECT COALESCE(MAX(line_no), 0) AS m FROM order_items WHERE order_id=?", (order_id,)
            ).fetchone()["m"]
            new_items = []
            for item in group["items"]:
                key = (item.get("material") or "", item.get("serial") or "")
                if existing_counts.get(key, 0) > 0:
                    existing_counts[key] -= 1
                    continue  # already on the board - additive-only, never rewritten
                new_items.append(item)
            for item in new_items:
                next_line += 1
                con.execute(
                    """
                    INSERT INTO order_items(order_id, line_no, description, material, serial, qty, mode, item_type)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        order_id, next_line, item["description"], item["material"], item["serial"],
                        item["qty"], item["mode"], item.get("item_type", ""),
                    ),
                )
            con.execute("UPDATE orders SET updated_at=? WHERE id=?", (now_text(), order_id))
            add_event(con, order_id, f"Excel 범위 불러오기({group['source_sheet']} {group['source_row_start']}-{group['source_row_end']}): 항목 {len(new_items)}개 추가")
            imported += 1
            continue
        if existing:
            order_id = existing["id"]
            # Real bug fixed 2026-08-03: this used to also overwrite
            # customer/phone/address/memo with whatever Excel
            # currently has, on EVERY reimport (the SAP auto-sync
            # loop runs this roughly every ~10 min for every order
            # still present in Excel) - silently reverting any
            # manual edit made via update_order_field() (the inline
            # cell editing added 2026-07-31) within one sync cycle,
            # with no visible error - it just looked like "my edit
            # didn't save" whenever the user happened to notice
            # later. item_date/board_status/deleted_at (per-item)
            # were already preserved across reimport; these four
            # order-level fields now get the same treatment - only
            # order_type and the source-sheet bookkeeping fields
            # (which aren't user-editable in workbench) still sync
            # from Excel every time.
            con.execute(
                """
                UPDATE orders SET order_type=?,
                    source_sheet=?, source_row_start=?, source_row_end=?, source_date=?, updated_at=?
                WHERE id=?
                """,
                (
                    group.get("order_type"), group.get("source_sheet"),
                    group.get("source_row_start"), group.get("source_row_end"), group_date,
                    now_text(), order_id,
                ),
            )
            prior_rows = con.execute(
                "SELECT id, line_no, description, material, serial, item_type, item_date, board_status, deleted_at "
                "FROM order_items WHERE order_id=?",
                (order_id,),
            ).fetchall()
            prior_ids = [prow["id"] for prow in prior_rows]
            # Which (item_id, field) pairs a human has directly
            # edited on the board - those fields must survive this
            # reimport untouched instead of taking Excel's fresh
            # value (see manual_item_edits' comment in ensure_db()
            # and its use just below in the insert loop).
            if prior_ids:
                placeholders = ",".join("?" for _ in prior_ids)
                for erow in con.execute(
                    f"SELECT item_id, field FROM manual_item_edits WHERE item_id IN ({placeholders})",
                    prior_ids,
                ).fetchall():
                    manual_edits_by_item.setdefault(erow["item_id"], set()).add(erow["field"])
            for prow in prior_rows:
                pkey = (prow["material"] or "", prow["serial"] or "")
                prior_by_key.setdefault(pkey, []).append(prow)
                prior_by_line[prow["line_no"]] = prow
            con.execute("DELETE FROM order_items WHERE order_id=?", (order_id,))
        else:
            cur = con.execute(
                """
                INSERT INTO orders(order_no, delivery_no, order_type, customer, phone, address, memo,
                    status, source_sheet, source_row_start, source_row_end, source_date, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, 'new', ?, ?, ?, ?, ?, ?)
                """,
                (
                    group.get("order_no"), group.get("delivery_no"), group.get("order_type"),
                    group.get("customer"), group.get("phone"), group.get("address"), group.get("memo"),
                    group.get("source_sheet"), group.get("source_row_start"), group.get("source_row_end"),
                    group_date, now_text(), now_text(),
                ),
            )
            order_id = cur.lastrowid
        for idx, item in enumerate(group["items"], start=1):
            prior = _pop_prior_item_state(prior_by_key, prior_by_line, item, idx)
            edited_fields = manual_edits_by_item.get(prior["id"], set()) if prior is not None else set()
            # A manually-edited field keeps its prior (human-typed)
            # value instead of taking Excel's fresh one - see
            # manual_item_edits' comment in ensure_db(). Fields not
            # in edited_fields still sync from Excel normally (e.g.
            # a serial SAP collects a few cycles after the item was
            # first created blank must still come through).
            description = prior["description"] if "description" in edited_fields else item["description"]
            material = prior["material"] if "material" in edited_fields else item["material"]
            serial = prior["serial"] if "serial" in edited_fields else item["serial"]
            # item_type added to this carry-over 2026-08-26, same reasoning as
            # description/material/serial above - a Type dropdown value the
            # user set by hand (now persisted via ITEM_EDIT_FIELDS, see its
            # own comment) shouldn't get silently overwritten by the next
            # emergency Excel reimport either.
            item_type = prior["item_type"] if prior is not None and "item_type" in edited_fields else item.get("item_type", "")
            cur = con.execute(
                """
                INSERT INTO order_items(order_id, line_no, description, material, serial, qty, mode, item_type)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    order_id, idx, description, material, serial,
                    item["qty"], item["mode"], item_type,
                ),
            )
            if prior is not None:
                if prior["item_date"] or prior["board_status"] or prior["deleted_at"]:
                    con.execute(
                        "UPDATE order_items SET item_date=?, board_status=?, deleted_at=? WHERE id=?",
                        (prior["item_date"], prior["board_status"], prior["deleted_at"], cur.lastrowid),
                    )
                # cell_highlights is keyed by item_id, which changes
                # on every reimport (order_items gets DELETEd and
                # reinserted above, not updated in place) - without
                # this, a highlight set on an item silently vanishes
                # the next time its order gets reimported. Re-point
                # it at the new id instead, same "carry state across
                # the delete+reinsert" treatment item_date/board_status/
                # deleted_at already get via prior above.
                con.execute(
                    "UPDATE cell_highlights SET item_id=? WHERE item_id=?",
                    (cur.lastrowid, prior["id"]),
                )
                if edited_fields:
                    con.execute(
                        "UPDATE manual_item_edits SET item_id=? WHERE item_id=?",
                        (cur.lastrowid, prior["id"]),
                    )
        add_event(con, order_id, f"Imported from Excel sheet {group['source_sheet']} rows {group['source_row_start']}-{group['source_row_end']}")
        imported += 1
    return imported


def load_today_from_excel():
    """Manual/emergency tool only - as of 2026-08-06, nothing calls this
    automatically anymore. SAP's own scripts (main.py/manual_order_handler.py)
    push freshly-collected orders straight into workbench.db via
    import_sap_rows() the moment they write them to Excel, instead of relying
    on this function to re-scan and re-parse the whole sheet ~every 10 min
    (that automatic Excel->Workbench round-trip was the actual source of the
    "workbench randomly goes down / edits vanish" reports - see
    project_sap_watchdog_fix memory - since every reimport had to re-derive
    structured data from Excel cell text and needed ever-growing per-field
    "don't overwrite this" carve-outs). Kept for real recovery scenarios
    (workbench.db lost/corrupted, or Excel was hand-edited and workbench
    needs to catch up) - use import_excel_range() instead when you only want
    a specific sheet/row range rather than all of today's sheet(s)."""
    ensure_db()
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    sheets = today_sheet_names(wb.sheetnames)
    if not sheets:
        wb.close()
        raise RuntimeError(f"today sheet not found in {EXCEL_PATH}")
    try:
        with connect() as con:
            imported = 0
            for sheet_name in sheets:
                groups = _scan_sheet_groups(wb[sheet_name], sheet_name)
                imported += _ingest_groups(con, groups)
    finally:
        wb.close()
    return {"imported": imported, "sheets": sheets}


def import_excel_range(sheet_name, row_start=1, row_end=None):
    """Manual, scoped emergency import - "이 시트의 이 행부터 이 행까지만
    가져오기", a one-shot, explicitly-targeted alternative to
    load_today_from_excel()'s whole-sheet scan. Meant for the rare case
    where something needs to be pulled back in from Excel by hand (e.g. a
    row someone typed directly into Excel, or recovering after a data
    problem) - NOT for routine syncing, which is now import_sap_rows()'s job
    (see load_today_from_excel()'s own docstring for why routine automatic
    Excel->Workbench reimport was removed)."""
    ensure_db()
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"sheet not found: {sheet_name}")
        ws = wb[sheet_name]
        row_start = max(1, int(row_start or 1))
        row_end = int(row_end) if row_end else ws.max_row
        groups = _scan_sheet_groups(ws, sheet_name, row_start=row_start, row_end=row_end)
        with connect() as con:
            # additive=True: this is a scoped/partial slice of the sheet by
            # design (see this function's own docstring) - an order caught
            # in that slice may have more items elsewhere that this scan
            # never saw, so existing orders must only ever gain items here,
            # never lose them (see _ingest_groups()'s docstring for the
            # real 2026-08-07 incident this fixes).
            imported = _ingest_groups(con, groups, additive=True)
    finally:
        wb.close()
    return {"imported": imported, "sheet": sheet_name, "rowStart": row_start, "rowEnd": row_end}


def import_sap_rows(rows):
    """The real, routine SAP->workbench path (replaces the old automatic
    load_today_from_excel() reimport, see its own docstring). Takes the exact
    same structured row shape excel_handler.write_orders_to_excel() consumes -
    one order's items, each a dict with order_prefix/order_type/order_num/
    obd/extra_orders/description/quantity/material/serial_number/customer/
    phone/company/street/street2/memo (see sap_handler.build_excel_rows(),
    zrma_handler.build_excel_rows_zrma(), vl10g_handler's row builder) -
    called by main.py/manual_order_handler.py right after they write the
    same rows to Excel, so workbench.db reflects a freshly-collected order
    without ever re-parsing Excel cell text for it.

    Deliberately additive-only, unlike the Excel-reimport path: finds-or-
    creates the orders row, then appends new order_items rows starting after
    whatever line_no already exists - it never deletes/rewrites existing
    items. Each SAP handler only calls this once per genuinely new order it
    just detected (order_tracker's processed-order tracking prevents the
    same order being processed twice), so there's no reimport-driven need for
    manual_item_edits-style field preservation here - a field just never gets
    overwritten because the row that already exists is never touched again."""
    if not rows:
        return {"imported": 0}
    first = rows[0]
    order_no = str(first.get("order_num") or "").strip()
    if not order_no:
        return {"imported": 0, "error": "missing order_num"}
    delivery_no = str(first.get("obd") or "").strip()
    order_type = first.get("order_type") or ""
    today_iso = date.today().isoformat()

    address_parts = []
    if first.get("company"):
        address_parts.append(first["company"])
    street, street2 = first.get("street") or "", first.get("street2") or ""
    if street and street2:
        address_parts.append(f"{street},\n{street2}")
    elif street:
        address_parts.append(street)
    elif street2:
        address_parts.append(street2)
    if first.get("cust_no"):
        address_parts.append(f"(cust# {first['cust_no']})")
    new_address = "\n".join(address_parts)

    ensure_db()
    with connect() as con:
        existing = con.execute(
            "SELECT id FROM orders WHERE order_no=? AND COALESCE(delivery_no,'')=COALESCE(?, '')",
            (order_no, delivery_no or ""),
        ).fetchone()
        if existing:
            order_id = existing["id"]
            # Real bug found 2026-08-07 (order 7837832, ZOR): the order's
            # very first collection attempt failed (ZZ block, RPC
            # disconnect) and got saved as a placeholder shell with
            # customer/phone blank and a "[Delivery Block: ...]" memo. A
            # later retry collected the real customer/phone/address/memo
            # fine, but this branch used to only bump updated_at - the
            # additive-only "never touch an existing order field again"
            # rule (meant to protect a human's manual edit) was also
            # silently locking in that first failed attempt's blank/
            # placeholder fields forever. Fix: still never overwrite a
            # field that already holds real data (human-edited or a prior
            # good collection), but DO backfill a field that's still
            # genuinely blank - there's nothing to protect there.
            cur_row = con.execute(
                "SELECT customer, phone, address, memo FROM orders WHERE id=?", (order_id,)
            ).fetchone()
            backfill = {
                "customer": first.get("customer") or "",
                "phone": first.get("phone") or "",
                "address": new_address,
                "memo": first.get("memo") or "",
            }
            set_clauses, params = [], []
            for field, new_val in backfill.items():
                if new_val and not (cur_row[field] or "").strip():
                    set_clauses.append(f"{field}=?")
                    params.append(new_val)
            set_clauses.append("updated_at=?")
            params.append(now_text())
            params.append(order_id)
            con.execute(f"UPDATE orders SET {', '.join(set_clauses)} WHERE id=?", params)
            max_line = con.execute(
                "SELECT COALESCE(MAX(line_no), 0) AS m FROM order_items WHERE order_id=?", (order_id,)
            ).fetchone()["m"]
        else:
            extra_codes = "\n".join(first.get("extra_orders") or [])
            cur = con.execute(
                """
                INSERT INTO orders(order_no, delivery_no, order_type, customer, phone, address, memo,
                    status, source_sheet, source_row_start, source_row_end, source_date, created_at, updated_at,
                    extra_codes)
                VALUES (?, ?, ?, ?, ?, ?, ?, 'new', ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    order_no, delivery_no, order_type, first.get("customer") or "", first.get("phone") or "",
                    new_address, first.get("memo") or "",
                    "SAP-direct", None, None, today_iso, now_text(), now_text(),
                    extra_codes,
                ),
            )
            order_id = cur.lastrowid
            max_line = 0

        for offset, item in enumerate(rows, start=1):
            serial = item.get("serial_number") or ""
            material = str(item.get("material") or "")
            try:
                qty = max(1, int(float(item.get("quantity") or 1)))
            except (ValueError, TypeError):
                qty = 1
            mode = "material_only" if (not serial and material and qty > 1) else "serial"
            item_type = item.get("order_prefix") or first.get("order_prefix") or ""
            con.execute(
                """
                INSERT INTO order_items(order_id, line_no, description, material, serial, qty, mode, item_type)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (order_id, max_line + offset, item.get("description") or "", material, serial, qty, mode, item_type),
            )
        add_event(con, order_id, f"SAP 자동수집 반영 ({len(rows)}행)")
    return {"imported": len(rows), "orderId": order_id}


EXPORT_HEADER = ["Order #", "item", "M/N", "S/N", "customer", "phone", "ADDRESS", "memo"]

_KOREAN_WEEKDAYS = ["월", "화", "수", "목", "금", "토", "일"]

# Column widths and per-column cell styling, matched against an existing
# hand-formatted sheet (verified against "7-27": header fill is the theme
# accent2 color tinted +0.8, or "FBE3D6" computed out to RGB; data-row fill
# is the plain blue "00B0F0" Excel uses for "done" rows). Column A carries
# its own item_type/order_label/delivery_no line on every row (never blank
# for continuation - the original sheets always repeat it), while columns
# E-H (customer/phone/address/memo) are merged down the whole customer
# group, same as the original sheets build via SAP collection.
_EXPORT_COL_WIDTHS = {1: 23.375, 2: 40.0, 3: 13.75, 4: 14.5, 5: 16.625, 6: 19.875, 7: 47.75, 8: 25.125}
_EXPORT_HEADER_FILL = PatternFill(fill_type="solid", fgColor="FBE3D6")
_EXPORT_DATA_FILL = PatternFill(fill_type="solid", fgColor="00B0F0")
_EXPORT_THIN_BORDER = Border(*(Side(style="thin"),) * 4)
_EXPORT_COL_ALIGN = {
    1: Alignment(vertical="top", wrap_text=True),
    2: Alignment(horizontal="left", vertical="center", wrap_text=True),
    3: Alignment(horizontal="center", vertical="center"),
    4: Alignment(horizontal="center", vertical="center"),
    5: Alignment(horizontal="center", vertical="center"),
    6: Alignment(horizontal="center", vertical="center", wrap_text=True),
    7: Alignment(horizontal="center", vertical="center", wrap_text=True),
    8: Alignment(horizontal="center", vertical="top", wrap_text=True),
}
# Manual cell_highlights carried into the export: column 1 (Order #) also
# carries the Type text (see _order_cell_text below, there's no separate
# Type column in this sheet layout - matches the hand-built originals), so a
# "type" highlight is checked as a fallback for column 1 if "order" isn't
# set. Every other column maps to exactly one highlight field.
_EXPORT_HIGHLIGHT_FIELDS_BY_COL = {
    1: ("order", "type"), 2: ("description",), 3: ("material",), 4: ("serial",),
    5: ("customer",), 6: ("phone",), 7: ("address",), 8: ("memo",),
}
_EXPORT_HIGHLIGHT_FILL = {
    "yellow": PatternFill(fill_type="solid", fgColor="FFFF00"),
    "pink": PatternFill(fill_type="solid", fgColor="FFC0CB"),
}
_EXPORT_HIGHLIGHT_FONT_COLOR = {"red": "FFFF0000"}


def _order_cell_text(row):
    order_cell = str(row.get("itemType") or "").strip()
    label = str(row.get("orderLabel") or "").strip()
    if label:
        order_cell = f"{order_cell} {label}".strip()
    codes = str(row.get("deliveryNo") or "").strip()
    if codes:
        order_cell = f"{order_cell}\n{codes}" if order_cell else codes
    return order_cell


def _write_export_sheet(wb, sheet_name, iso, rows):
    """Writes `rows` into `sheet_name`. If that sheet doesn't exist yet,
    creates it fresh (date banner + headers + data, as before). If it
    already exists - a prior 마감 for this exact date already made it -
    appends `rows` starting right after its last used row instead: no new
    "(2)"-suffixed sheet, no rewriting the banner/header above.

    2026-08-27 (user request): an order whose 마감 slips past its own day
    (e.g. a POD/SAP issue only gets resolved a few days later) should stay on
    the workbench board, unchecked, while everything else for that date gets
    마감-ed normally; once it's later resolved and checked off, 마감-ing it
    should land it at the bottom of the date it always belonged to - not
    spawn a duplicate sheet for that date. Since export_dates_to_excel() now
    always targets the exact "{month}-{day}" sheet name (no more
    _unique_sheet_name() collision-avoidance), this is the only place that
    needs to branch on "does it already exist"."""
    d = date.fromisoformat(iso)
    existing = sheet_name in wb.sheetnames
    ws = wb[sheet_name] if existing else wb.create_sheet(sheet_name)

    if existing:
        row_cursor = max(ws.max_row + 1, 3)
    else:
        for col, width in _EXPORT_COL_WIDTHS.items():
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

        # Row 1: date banner (bold, only the ADDRESS column carries the text,
        # same position as the original sheets).
        date_cell = ws.cell(row=1, column=7)
        date_cell.value = f"{d.year}년 {d.month}월 {d.day}일 {_KOREAN_WEEKDAYS[d.weekday()]}요일"
        date_cell.font = Font(bold=True)
        date_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Row 2: column headers.
        for col, label in enumerate(EXPORT_HEADER, start=1):
            cell = ws.cell(row=2, column=col)
            cell.value = label
            cell.fill = _EXPORT_HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = _EXPORT_THIN_BORDER
        row_cursor = 3

    # Data rows. Each of E-H (customer/phone/address/memo) merges down its
    # OWN contiguous run independently, per row["mergeFirst"][field] (built
    # client-side by buildExportRows/computeFieldSpans) - a real case can
    # share e.g. just the address across rows with different customer names,
    # so this can no longer be one shared group boundary for all 4 columns.
    # This merge run is scoped only to the rows being written THIS call - an
    # appended block never merges upward into whatever's already above it.
    n = len(rows)
    row_numbers = [row_cursor + i for i in range(n)]

    for i, row in enumerate(rows):
        r = row_numbers[i]
        mf = row.get("mergeFirst") or {}
        values = {
            1: _order_cell_text(row),
            2: row.get("description") or "",
            3: row.get("material") or "",
            4: row.get("serial") or "",
            5: (row.get("customer") or "") if mf.get("customer", True) else "",
            6: (row.get("phone") or "") if mf.get("phone", True) else "",
            7: (row.get("address") or "") if mf.get("address", True) else "",
            8: (row.get("memo") or "") if mf.get("memo", True) else "",
        }
        highlights = row.get("highlights") or {}
        for col in range(1, 9):
            cell = ws.cell(row=r, column=col)
            cell.value = values[col]
            style = next(
                (highlights[f] for f in _EXPORT_HIGHLIGHT_FIELDS_BY_COL[col] if highlights.get(f)), None
            )
            cell.fill = _EXPORT_HIGHLIGHT_FILL.get((style or {}).get("bg"), _EXPORT_DATA_FILL)
            font_kwargs = {}
            if col == 8:
                font_kwargs["bold"] = True
            if style and style.get("bold"):
                font_kwargs["bold"] = True
            if style and style.get("color") in _EXPORT_HIGHLIGHT_FONT_COLOR:
                font_kwargs["color"] = _EXPORT_HIGHLIGHT_FONT_COLOR[style["color"]]
            if font_kwargs:
                cell.font = Font(**font_kwargs)
            cell.alignment = _EXPORT_COL_ALIGN[col]
            cell.border = _EXPORT_THIN_BORDER

    for field, col in zip(("customer", "phone", "address", "memo"), (5, 6, 7, 8)):
        i = 0
        while i < n:
            j = i + 1
            while j < n and not (rows[j].get("mergeFirst") or {}).get(field, True):
                j += 1
            if j - i > 1:
                ws.merge_cells(start_row=row_numbers[i], start_column=col, end_row=row_numbers[j - 1], end_column=col)
            i = j

    row_cursor = (row_numbers[-1] + 1) if n else row_cursor
    return ws, row_cursor - 1


def export_dates_to_excel(dates_payload):
    """
    Writes the board's current (client-side) state for the given dates into
    the "마감" export workbook (FINALIZE_EXPORT_PATH - a dedicated file,
    separate from the live daily-collection workbook EXCEL_PATH, per user
    request 2026-08-14) - one sheet per date, named the same "{month}-{day}"
    way `today_sheet_names()` looks for them (e.g. "7-29"). Existing sheets
    are never deleted or overwritten: a date exported for the first time
    gets a fresh sheet, but 마감-ing that SAME date again later (e.g. one
    order's 마감 slipped a few days past everything else - 2026-08-27 user
    request) appends the newly-checked rows to the bottom of that same
    "7-29" sheet instead of spawning "7-29 (2)" - see _write_export_sheet().
    The board's data still comes from the client's in-memory `ROWS`, not
    workbench.db, since nothing this session persists there yet. Formatting
    (column widths, header fill, blue "done" fill, borders, merged
    customer/phone/address/memo per group) matches the hand-built sheets SAP
    collection produces.

    Takes a full backup of the export workbook before writing - this is the
    one action in this app that touches a real, persistent Excel file
    (everything else here is screen-only), so a mistake here is not casually
    undoable the way an in-browser action is.

    Once the save below succeeds, also kicks off vendor_dashboard.py in the
    background for every exported date (2026-08-25: 마감 IS what freezes a
    day's colors/rows into this same file, so it's the natural point to keep
    the local Vendor Activities Dashboard copy/paste sheet in sync - no
    separate manual step to remember). That refresh is fire-and-forget: its
    own failure must never make 마감 itself look like it failed.
    """
    excel_path = Path(FINALIZE_EXPORT_PATH)
    if not excel_path.exists():
        # openpyxl은 시트가 0개인 워크북 저장을 거부하므로 기본 빈 "Sheet"는
        # 그대로 둔다 - 실제 마감 시트들과 나란히 있어도 무해함.
        excel_path.parent.mkdir(parents=True, exist_ok=True)
        openpyxl.Workbook().save(excel_path)

    # 마감을 누를 때마다 새 타임스탬프 파일을 만들면 이 파일 자체가 이미
    # "마감 전용 백업 파일"이라 backup-of-a-backup이 무한정 쌓인다 (사용자
    # 지적, 2026-08-14). 그래서 백업은 이름 고정 1개만 두고 매번 덮어쓴다 -
    # "직전 저장 시점으로 한 단계만 되돌릴 수 있는" 안전장치는 유지하되
    # 파일이 쌓이지는 않음.
    backup_path = excel_path.with_name(f"{excel_path.stem}.before_last_export{excel_path.suffix}")
    shutil.copy2(excel_path, backup_path)

    wb = openpyxl.load_workbook(excel_path)
    try:
        result = {}
        for iso, rows in dates_payload.items():
            d = date.fromisoformat(iso)
            sheet_name = f"{d.month}-{d.day}"
            _write_export_sheet(wb, sheet_name, iso, rows)
            result[iso] = {"sheet": sheet_name, "rows": len(rows)}
        wb.save(excel_path)
        _trigger_dashboard_refresh(list(dates_payload.keys()))
        _mirror_finalize_to_shared_excel(dates_payload)
        return {"exported": result, "backup": str(backup_path)}
    finally:
        wb.close()


def _mirror_finalize_to_shared_excel(dates_payload):
    """마감(확정)된 데이터를 공유파일(EXCEL_PATH)에도 그대로 반영한다
    (2026-09-01 사용자 결정).

    배경: 공유파일은 다른 사람들과 공유하는 파일이라 항상 깔끔해야 하는데,
    지금까지는 매일 SAP 자동수집이 "아직 확정 전"인 진행중 상태까지 직접
    써서 지저분해졌었다. 그 자동쓰기는 이제 껐고(main.py/manual_order_
    handler.py), 대신 workbench에서 실제로 "완료"로 확정 처리된 마감 데이터만
    이 함수로 공유파일에 반영한다 - 공유파일이 "1년치 확정 기록 아카이브"
    역할을 하게 됨. 실시간 상태는 board_sync.py가 20분마다 자동으로 채우는
    REALTIME_SYNC_PATH가 대신한다.

    구조 문제와 해결: 공유파일의 기존 날짜별 시트("M-D")는 이 마감 파일과
    달리 한 시트 안에 여러 날짜 배너가 쌓이는 옛 구조(board_sync.py 참고)라,
    "한 시트=한 날짜" 가정으로 그대로 이어붙이면 데이터가 시트 맨 끝(전혀
    다른 미래 날짜 배너 밑)에 엉뚱하게 붙는다. 그런 시트(배너가 2개 이상)를
    만나면 건드리지 않고 "{sheet_name} (마감)"이라는 구분되는 이름으로 대신
    만든다(excel_handler.mirror_finalized_rows_to_excel_path 안에서 처리) -
    앞으로 새로 마감되는 날짜는 이 옛 구조 시트가 없을 것이므로(이 전환
    이후로는 공유파일에 새 날짜 시트가 자동생성 안 됨) 자연스럽게 깨끗한
    "한 시트=한 날짜" 구조로 자리잡는다.

    2026-09-02 openpyxl→COM 전환: 처음엔 openpyxl(_write_export_sheet 재사용
    시도)로 짰다가 실제로 두 번 다 PermissionError로 실패했다 - 공유파일은
    사람이 실제로 Excel에 열어놓고 보는 경우가 흔한데, openpyxl의 일반 파일
    저장은 그렇게 열려있는 파일에 못 쓴다. write_orders_to_excel()이 이미
    검증해온 xlwings/COM 경로(excel_handler.get_workbook())로 옮겼다 - 열려
    있는 세션에 그대로 붙어서 실시간 반영된다.

    실패해도 마감 자체(위에서 이미 저장 완료)를 막지 않는다 - _trigger_
    dashboard_refresh()와 같은 이유로 로그만 남기고 넘어간다."""
    excel_path = Path(EXCEL_PATH)
    if not excel_path.exists():
        logger.warning(f"공유파일 없음, 마감 반영 스킵: {excel_path}")
        return

    backup_path = excel_path.with_name(f"{excel_path.stem}.before_last_마감{excel_path.suffix}")
    try:
        shutil.copy2(excel_path, backup_path)
    except Exception:
        logger.error("공유파일 백업 실패 - 반영은 계속 시도", exc_info=True)

    try:
        excel_handler.mirror_finalized_rows_to_excel_path(dates_payload)
        logger.info(f"공유파일 마감 반영 완료: {list(dates_payload.keys())}")
    except Exception:
        logger.error("공유파일(EXCEL_PATH) 마감 반영 실패 (마감 자체는 정상 완료됨)", exc_info=True)


def _trigger_dashboard_refresh(iso_dates):
    """Fire-and-forget: launch vendor_dashboard.py in the background to
    refresh the local Vendor Activities Dashboard copy/paste sheet for the
    dates just 마감-ed. No console window (same pattern as every other
    background launch in this app - see startup.py/main.py above), and any
    failure to even start it is logged and swallowed rather than raised,
    since the 마감 export above has already succeeded by this point."""
    if not iso_dates:
        return
    try:
        subprocess.Popen(
            [sys.executable, "vendor_dashboard.py", "--dates", ",".join(iso_dates)],
            cwd=str(BASE_DIR),
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception:
        logger.warning("Dashboard 자동 갱신 실행 실패 (마감 자체는 정상 완료됨)", exc_info=True)


def sync_delivery_excel():
    """workbench.db의 현재 상태를 기존 배송장 엑셀(EXCEL_PATH - 예전부터 써온
    실시간 배송장, board_sync.py 상단 docstring 참고)의 오늘 시트에 반영한다.
    board_sync.sync()를 그 자리에서 바로 호출(서브프로세스 아님 - 순수
    파이썬+xlwings 호출이라 콘솔 창 걱정이 없고, 결과 건수를 그대로 돌려줄
    수 있어 마감 버튼처럼 toast에 몇 건 반영됐는지 보여줄 수 있음).

    2026-08-28 사용자 요청: workbench가 진짜 메인이 될 때까지, 예전부터
    써온 배송장 엑셀도 손으로 안 맞춰도 되게 하기 위한 임시 다리 -
    workbench가 충분히 안정화되면 이 버튼째로 치우면 됨. 이미 있는 오더는
    건드리지 않고 누락된 것만 추가하므로 여러 번 눌러도 안전."""
    import board_sync
    return board_sync.sync()


