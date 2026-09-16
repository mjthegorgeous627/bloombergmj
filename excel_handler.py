"""
Excel 핸들러 - xlwings 사용으로 Excel 열린 상태에서도 실시간 데이터 입력 가능.
"""

import time
import xlwings as xw
import openpyxl
import re
from datetime import datetime
from config import EXCEL_PATH
from order_lock import ExcelWriteLock

_DATE_RE = re.compile(r'(\d{1,2})월\s*(\d{1,2})일')


def _com_retry(fn, attempts=3, delay=1.0):
    """xlwings/COM 호출(Insert/Merge/used_range 등)은 일시적 원인으로 실패했다가
    그냥 다시 하면 되는 경우가 실제로 많다 - automation.log에 5개월간 12번 기록된
    OLE error 0x800ac472 / RPC 예외 -2147352567가 전부 이 유형이었다(2026-08-24
    구조 점검). ExcelWriteLock으로 프로세스 간 동시 접근이라는 주된 원인은
    막았지만, Excel이 그 순간 다른 이유로 잠깐 바쁜 경우까지는 배제할 수 없어
    방어선을 하나 더 둔다. 마지막 시도까지 실패하면 원래 예외를 그대로 올려서,
    호출부의 기존 except 처리가 지금처럼 "이 오더만 실패, 다음 사이클 재시도"로
    다루게 한다 - 동작을 안 바꾸고 성공률만 높이는 게 목적."""
    last_exc = None
    for i in range(attempts):
        try:
            return fn()
        except Exception as e:
            last_exc = e
            if i < attempts - 1:
                time.sleep(delay)
    raise last_exc


def get_today_sheet_name():
    """시트명 형식: '3-17' (월-일)."""
    today = datetime.today()
    return f"{today.month}-{today.day}"


def get_workbook():
    """
    이미 열려있는 Excel 파일에 연결. 없으면 파일 열기.
    반환: xlwings Book 객체

    2026-09-01 버그 수정: 예전엔 xw.books(활성 Excel 인스턴스 하나의 책 목록)를
    파일명만("book.name")으로 훑었는데, 다른 폴더에 같은 파일명의 별개 파일이
    동시에 열려 있으면(실제로 이날 C:\\Users\\bloomberg\\Documents\\에 동명 사본이
    떠 있어서 발생) 엉뚱한 파일을 골라 거기에 쓰고 저장해버리는 사고가 났다
    (진짜 EXCEL_PATH 파일은 안 건드려진 채로 남음). xw.Book(EXCEL_PATH)은 전체
    경로(fullname) 기준으로, 떠 있는 모든 Excel 인스턴스를 통틀어 정확히 같은
    파일을 찾아 붙거나(없으면 새로 연다) 동작이 확인돼(실측) 그걸로 대체한다.
    """
    return xw.Book(EXCEL_PATH)


def find_today_section_end(ws):
    """
    오늘 날짜 섹션의 마지막 데이터 행 번호 반환.
    - 날짜 헤더: G열에 '3월 19일', '3월19일 목요일', '2026년 3월 19일' 등
    - 오후 섹션: '3월 23일 오후' 처럼 오후 명시 시 해당 섹션을 우선 사용
    - 'Client/Bloomberg - Scheduled' 병합 행 이전까지만 스캔
    """
    today = datetime.today()
    today_m = today.month
    today_d = today.day

    used = _com_retry(lambda: ws.used_range)
    max_row = used.last_cell.row
    max_col = 8

    # 'Scheduled' 병합 행 경계 먼저 탐지 (Client/Bloomberg scheduled 구분선)
    scheduled_boundary = max_row + 1
    for r in range(1, max_row + 1):
        try:
            if ws.cells(r, 1).api.MergeCells:
                val = ws.cells(r, 1).value
                if val and isinstance(val, str) and 'scheduled' in val.lower():
                    scheduled_boundary = r
                    break
        except Exception:
            pass

    # G열(7열)에서 날짜 헤더 행 수집 (scheduled 경계 이전만)
    date_headers = []  # [(row, month, day), ...]
    for r in range(1, scheduled_boundary):
        val = ws.cells(r, 7).value
        if val and isinstance(val, str):
            m = _DATE_RE.search(val)
            if m:
                date_headers.append((r, int(m.group(1)), int(m.group(2))))

    # 오늘 날짜 헤더 찾기 (마지막 매칭 사용 → 오후 섹션이 있으면 오후 우선)
    today_header_row = None
    today_idx = None
    for i, (r, mo, d) in enumerate(date_headers):
        if mo == today_m and d == today_d:
            today_header_row = r
            today_idx = i  # break하지 않음 → 마지막 매칭 사용

    if today_header_row is None:
        return scheduled_boundary - 1

    # 다음 경계: 다른 날짜 헤더 또는 scheduled 경계
    next_boundary = scheduled_boundary
    for j in range(today_idx + 1, len(date_headers)):
        nr, nmo, nd = date_headers[j]
        if nmo != today_m or nd != today_d:
            next_boundary = nr
            break

    # 오늘 섹션의 마지막 데이터 행 (연속 빈 행 3개 이상이면 종료)
    last_data_row = today_header_row + 1
    consecutive_empty = 0
    for r in range(today_header_row + 2, next_boundary):
        if any(ws.cells(r, c).value for c in range(1, max_col + 1)):
            last_data_row = r
            consecutive_empty = 0
        else:
            consecutive_empty += 1
            if consecutive_empty >= 3:
                break

    return last_data_row


def write_orders_to_excel(order_data_list):
    """오더 번호와 무관하게 실제 Excel COM 쓰기 구간 전체를 전역 직렬화 -
    ExcelWriteLock 설명(order_lock.py) 참고. 실제 작업은 그대로
    _write_orders_to_excel_locked()가 한다."""
    with ExcelWriteLock():
        return _write_orders_to_excel_locked(order_data_list)


def _write_orders_to_excel_locked(order_data_list):
    """
    order_data_list: 한 오더의 아이템들 리스트.
    Excel이 열려있어도 실시간으로 반영됨.
    """
    wb = get_workbook()

    # 시트 찾기
    sheet_name = get_today_sheet_name()
    sheet_names = [s.name for s in wb.sheets]

    if sheet_name in sheet_names:
        ws = wb.sheets[sheet_name]
    else:
        # 정확한 "월-일" 시트명이 없을 때, "월.일"/"월_일" 같은 구분자
        # 변형만 추가로 시도한다 (db.py의 today_sheet_names()와 동일한
        # 후보군). 예전엔 시트명에 오늘 월/일 숫자가 부분 문자열로만
        # 들어있어도 매치시켰는데(예: 1월 3일에 "1-13"/"11-3"/"1-31" 등을
        # 전부 오매칭), 그마저 안 맞으면 wb.sheets.active(=사용자가 지금
        # 실제로 보고 있는, 완전히 무관할 수 있는 시트)에 새 오더 행을
        # 그대로 꽂아넣었다 - 조용한 오분류/데이터 오염 위험이라, 못 찾으면
        # 화면 서식을 바꾸는 대신 명확히 실패시킨다.
        today = datetime.today()
        candidates = {
            f"{today.month}-{today.day}",
            f"{today.month}.{today.day}",
            f"{today.month}_{today.day}",
        }
        matched = next((s for s in sheet_names if s in candidates), None)
        if matched is None:
            raise RuntimeError(
                f"오늘({sheet_name}) 시트를 찾을 수 없습니다 - 시트가 아직 "
                f"안 만들어졌거나 이름이 다릅니다. 기존 시트 목록: {sheet_names}"
            )
        ws = wb.sheets[matched]

    # 삽입 위치 결정
    insert_after_row = find_today_section_end(ws)
    insert_at = insert_after_row + 1
    num_rows = len(order_data_list)

    # 행 삽입: 한 행씩 삽입 (안정적)
    for _ in range(num_rows):
        _com_retry(lambda: ws.range(f"A{insert_at}:H{insert_at}").api.EntireRow.Insert())

    # 삽입된 행 배경색 흰색으로 초기화 (이전 행 서식 상속 방지)
    ws.range(ws.cells(insert_at, 1), ws.cells(insert_at + num_rows - 1, 8)).color = (255, 255, 255)

    # 데이터 입력
    for idx, item in enumerate(order_data_list):
        row_num = insert_at + idx

        # A열: Order # (배송/회수 prefix + 오더번호 + 추가 오더번호들)
        order_line = f"{item['order_prefix']} {item['order_type']} {item['order_num']}"
        if item.get('obd'):
            order_line += f"\nOBD {item['obd']}"
        if item.get('extra_orders'):
            for eo in item['extra_orders']:
                order_line += f"\n{eo}"
        cell_a = ws.cells(row_num, 1)
        cell_a.value = order_line
        cell_a.api.WrapText = True

        # B열: item
        description = item.get('description', '')
        try:
            qty_value = int(float(item.get('quantity') or 1))
        except (ValueError, TypeError):
            qty_value = 1
        if qty_value > 1:
            description = f"{description}\nQty: {qty_value}" if description else f"Qty: {qty_value}"
        cell_b = ws.cells(row_num, 2)
        cell_b.value = description
        cell_b.api.WrapText = True

        # C열: M/N (숫자형 방지 - 문자열로 저장)
        ws.cells(row_num, 3).value = str(item.get('material', ''))

        # D열: S/N (회수 아이템은 SAP에서 수집, 없으면 X, 배송은 공백)
        ws.cells(row_num, 4).value = item.get('serial_number', '')

        # E열: 담당자
        ws.cells(row_num, 5).value = item.get('customer', '')

        # F열: 전화번호
        ws.cells(row_num, 6).value = item.get('phone', '')

        # G열: 주소 (모든 행에 기입)
        address_parts = []
        if item.get('company'):
            address_parts.append(item['company'])
        street = item.get('street', '')
        street2 = item.get('street2', '')
        if street and street2:
            address_parts.append(f"{street},\n{street2}")
        elif street:
            address_parts.append(street)
        elif street2:
            address_parts.append(street2)
        # Ship-to Party's own SAP number (Cust#) - appended as its own
        # trailing line, not a separate column (user's call, 2026-08-06).
        if item.get('cust_no'):
            address_parts.append(f"(cust# {item['cust_no']})")
        cell_g = ws.cells(row_num, 7)
        cell_g.value = '\n'.join(address_parts)
        cell_g.api.WrapText = True

        # H열: 메모 (모든 행에 기입). 수집 시각 타임스탬프는 2026-08-06부로 안 붙임 -
        # 사용자가 실제 배송/회수 관련 날짜(delivery date, coll due date 등)와
        # 헷갈린다고 해서 제거.
        memo = item.get('memo', '')
        cell_h = ws.cells(row_num, 8)
        cell_h.value = memo
        cell_h.api.WrapText = True

    # E/F/G/H 열 병합 (같은 오더의 여러 행)
    #
    # 2026-08-31 실사고(오더 67084542, ZRX): 새로 삽입한 행이 바로 위의 기존
    # 병합 블록(이전 오더의 E~H 병합)에 바로 붙어서 들어가면, EntireRow.Insert()가
    # 그 인접 병합의 서식을 새 행에도 일부 상속시켜 놓는 경우가 있다 - 새 범위가
    # "이미 부분적으로/다르게 병합된" 상태로 시작하는 셈이라, 그 위에 Merge()를
    # 부르면 Excel이 설명 없는 일반 COM 오류(0x800A03EC)로 거부한다. _com_retry의
    # 3회 재시도로는 구조적 충돌이라 절대 안 풀려서(실측: 같은 오더로 두 번 재현,
    # 매번 3번 다 실패) 매번 이 오더만 통째로 미수집 상태로 남았다 - 그 전까지는
    # "가끔 실패하는 COM 오류"로 오인해 재시도만으로 버텨온 걸로 보인다(automation.log
    # 5개월 12회 기록 중 일부가 실은 이 케이스였을 가능성). Merge() 직전에
    # UnMerge()로 상속된 잔여 병합 상태를 먼저 지워 항상 깨끗한 범위에서 병합되게
    # 한다 - 병합 안 된 범위에 UnMerge()를 불러도 안전한 무동작이라 부작용 없음.
    if num_rows > 1:
        end_row = insert_at + num_rows - 1
        wb.app.display_alerts = False

        def _merge_col(col):
            rng = ws.range(ws.cells(insert_at, col), ws.cells(end_row, col))
            rng.api.UnMerge()
            rng.api.Merge()

        try:
            for col in [5, 6, 7, 8]:  # E=담당자, F=전화, G=주소, H=메모
                _com_retry(lambda col=col: _merge_col(col))
        finally:
            # _com_retry가 3번 다 실패해서 예외를 그대로 올리는 경우(위 주석의
            # 2026-08-31 실사고 케이스), finally 없이는 display_alerts=False가
            # 이 공유 Excel App 인스턴스에 그대로 남아 이후 모든 작업(자동화는
            # 물론 사용자가 직접 여는 저장/덮어쓰기 확인창까지) 조용히 억제됨.
            wb.app.display_alerts = True

    # 모든 테두리 적용 (A~H, insert_at ~ insert_at+num_rows-1)
    border_range = ws.range(
        ws.cells(insert_at, 1),
        ws.cells(insert_at + num_rows - 1, 8)
    )
    for border_idx in [7, 8, 9, 10, 11, 12]:  # Left/Top/Bottom/Right/InsideV/InsideH
        border_range.api.Borders(border_idx).LineStyle = 1   # xlContinuous
        border_range.api.Borders(border_idx).Weight = 2      # xlThin

    # 저장
    wb.save()
    end_row = insert_at + num_rows - 1
    print(f"[Excel] {num_rows}개 행 삽입 완료 → {ws.name} 시트 행 {insert_at}~{end_row}")
    return ws, insert_at, end_row


def write_kakao_sent(ws, start_row, end_row):
    """write_orders_to_excel()과 같은 이유로 전역 락 적용 - 실제 작업은
    _write_kakao_sent_locked()가 한다."""
    with ExcelWriteLock():
        return _write_kakao_sent_locked(ws, start_row, end_row)


def _write_kakao_sent_locked(ws, start_row, end_row):
    """
    J열에 카톡 전송 시각 기록.
    여러 행이면 병합해서 한 칸에 표시.
    프린트 범위(A~H)에 포함되지 않아 인쇄에 영향 없음.
    """
    timestamp = datetime.now().strftime('카톡 ✓ %y.%m.%d %H:%M')
    wb = ws.book
    if end_row > start_row:
        wb.app.display_alerts = False
        try:
            _com_retry(lambda: ws.range(ws.cells(start_row, 10), ws.cells(end_row, 10)).api.Merge())
        finally:
            wb.app.display_alerts = True
    cell = ws.cells(start_row, 10)
    cell.value = timestamp
    cell.api.WrapText = True
    wb.save()
    print(f"[Excel] J열 카톡 전송 기록: {timestamp}")


def _is_recent_sheet(sheet_name, days=90):
    """시트명(M-D 형식)이 최근 N일 내인지 확인."""
    try:
        m, d = map(int, sheet_name.split('-'))
        today = datetime.today()
        sheet_date = datetime(today.year, m, d)
        if sheet_date > today:
            sheet_date = datetime(today.year - 1, m, d)
        return (today - sheet_date).days <= days
    except Exception:
        return False


def get_existing_order_numbers(days=45):
    """
    Excel 파일의 A열(Order #)에서 이미 입력된 오더번호(7자리 이상 숫자) 추출.
    최근 N일 시트만 스캔하여 성능 최적화.
    """
    found = set()
    num_re = re.compile(r'\b(\d{7,})\b')
    wb = None
    try:
        wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
        for sname in wb.sheetnames:
            if not _is_recent_sheet(sname, days=days):
                continue
            ws = wb[sname]
            for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
                val = row[0]
                if val and isinstance(val, str):
                    for m in num_re.finditer(val):
                        found.add(m.group(1))
    except Exception as e:
        print(f"[Excel] 기존 오더 스캔 실패: {e}")
    finally:
        if wb is not None:
            wb.close()
    return found


# ── 마감(확정) 데이터를 공유파일(EXCEL_PATH)에 반영 (2026-09-02) ─────────────
# workbench_app.py의 _write_export_sheet()(openpyxl)와 필드/서식은 거의
# 동일하지만, 이 파일은 사람이 실제로 열어놓고 보는 경우가 흔해서(2026-09-01
# 실사용에서 PermissionError로 두 번 실패 확인됨 - openpyxl은 이미 열려있는
# Excel과 같은 파일에 저장을 못 함) openpyxl 대신 xlwings/COM으로 쓴다 -
# get_workbook()처럼 이미 열린 세션에 붙어 그대로 실시간 반영된다.

_FINALIZE_MIRROR_COL_WIDTHS = {1: 23.375, 2: 40.0, 3: 13.75, 4: 14.5, 5: 16.625, 6: 19.875, 7: 47.75, 8: 25.125}
_FINALIZE_MIRROR_HEADER = ["Order #", "item", "M/N", "S/N", "customer", "phone", "ADDRESS", "memo"]
_FINALIZE_MIRROR_HEADER_RGB = (251, 227, 214)  # FBE3D6
# 마감 대상은 전부 완료(파랑) 처리된 것들이므로, board_sync.py가 오늘 도입한
# 상태색(done=#bfdbfe)과 통일 - 실시간동기화/공유파일이 같은 색 언어를 쓰게 함.
_FINALIZE_MIRROR_DATA_RGB = (191, 219, 254)
_FINALIZE_MIRROR_HL_BG = {"yellow": (255, 255, 0), "pink": (255, 192, 203)}
_FINALIZE_MIRROR_HL_FONT = {"red": (255, 0, 0)}
_FINALIZE_MIRROR_HL_FIELDS_BY_COL = {
    1: ("order", "type"), 2: ("description",), 3: ("material",), 4: ("serial",),
    5: ("customer",), 6: ("phone",), 7: ("address",), 8: ("memo",),
}
_WEEKDAYS_KO = ["월", "화", "수", "목", "금", "토", "일"]


def _finalize_mirror_order_cell_text(row):
    order_cell = str(row.get("itemType") or "").strip()
    label = str(row.get("orderLabel") or "").strip()
    if label:
        order_cell = f"{order_cell} {label}".strip()
    codes = str(row.get("deliveryNo") or "").strip()
    if codes:
        order_cell = f"{order_cell}\n{codes}" if order_cell else codes
    return order_cell


def write_finalized_sheet_com(wb, sheet_name, d, rows):
    """마감(확정)된 rows(workbench_app.py의 _write_export_sheet가 받는 것과
    동일한 dict 모양 - itemType/orderLabel/deliveryNo/description/material/
    serial/customer/phone/address/memo/mergeFirst/highlights)를 wb(이미 열려
    있는 EXCEL_PATH의 xlwings Book)의 sheet_name 시트에 COM으로 쓴다.
    시트가 없으면 새로 만들고(배너+헤더+열너비), 있으면 맨 끝에 이어붙인다
    (_write_export_sheet와 동일한 정책 - 같은 날짜가 나중에 또 마감되면
    이어붙임)."""
    names = [s.name for s in wb.sheets]
    existing = sheet_name in names
    if existing:
        ws = wb.sheets[sheet_name]
        used = _com_retry(lambda: ws.used_range)
        row_cursor = max(used.last_cell.row + 1, 3)
    else:
        ws = wb.sheets.add(sheet_name, after=wb.sheets[-1])
        for col, width in _FINALIZE_MIRROR_COL_WIDTHS.items():
            _com_retry(lambda col=col, width=width: setattr(ws.api.Columns(col), "ColumnWidth", width))
        banner_cell = ws.cells(1, 7)
        banner_cell.value = f"{d.year}년 {d.month}월 {d.day}일 {_WEEKDAYS_KO[d.weekday()]}요일"
        banner_cell.api.Font.Bold = True
        banner_cell.api.WrapText = True
        for col, label in enumerate(_FINALIZE_MIRROR_HEADER, start=1):
            c = ws.cells(2, col)
            c.value = label
            c.color = _FINALIZE_MIRROR_HEADER_RGB
        row_cursor = 3

    n = len(rows)
    if n == 0:
        return ws
    row_numbers = [row_cursor + i for i in range(n)]

    for i, row in enumerate(rows):
        r = row_numbers[i]
        mf = row.get("mergeFirst") or {}
        values = {
            1: _finalize_mirror_order_cell_text(row),
            2: row.get("description") or "",
            3: row.get("material") or "",
            4: row.get("serial") or "",
            5: (row.get("customer") or "") if mf.get("customer", True) else "",
            6: (row.get("phone") or "") if mf.get("phone", True) else "",
            7: (row.get("address") or "") if mf.get("address", True) else "",
            8: (row.get("memo") or "") if mf.get("memo", True) else "",
        }
        highlights = row.get("highlights") or {}
        for col in range(1, 9):
            cell = ws.cells(r, col)
            cell.value = values[col]
            style = next(
                (highlights[f] for f in _FINALIZE_MIRROR_HL_FIELDS_BY_COL[col] if highlights.get(f)), None
            )
            cell.color = _FINALIZE_MIRROR_HL_BG.get((style or {}).get("bg"), _FINALIZE_MIRROR_DATA_RGB)
            if col == 8 or (style and style.get("bold")):
                cell.api.Font.Bold = True
            if style and style.get("color") in _FINALIZE_MIRROR_HL_FONT:
                cell.font.color = _FINALIZE_MIRROR_HL_FONT[style["color"]]
            if col in (1, 2, 6, 7, 8):
                cell.api.WrapText = True

        border_range = ws.range(ws.cells(r, 1), ws.cells(r, 8))
        for idx in [7, 8, 9, 10, 11, 12]:
            border_range.api.Borders(idx).LineStyle = 1

    for field, col in zip(("customer", "phone", "address", "memo"), (5, 6, 7, 8)):
        i = 0
        while i < n:
            j = i + 1
            while j < n and not (rows[j].get("mergeFirst") or {}).get(field, True):
                j += 1
            if j - i > 1:
                wb.app.display_alerts = False
                try:
                    _com_retry(lambda i=i, j=j, col=col: ws.range(
                        ws.cells(row_numbers[i], col), ws.cells(row_numbers[j - 1], col)
                    ).api.Merge())
                finally:
                    wb.app.display_alerts = True
            i = j

    _com_retry(lambda: ws.api.Rows(f"{row_cursor}:{row_numbers[-1]}").AutoFit())
    return ws


def mirror_finalized_rows_to_excel_path(dates_payload):
    """마감(확정)된 dates_payload({iso날짜: rows})를 공유파일(EXCEL_PATH)에
    반영한다 - workbench_app.py의 export_dates_to_excel()이 성공적으로 저장한
    직후 호출됨(그쪽 마감 자체와는 독립, 실패해도 마감을 막지 않음 - 호출부가
    책임짐).

    공유파일의 기존 "M-D" 시트는 board_sync.py 방식(한 시트에 여러 날짜
    배너가 쌓이는 구조)일 수 있는데, 그런 시트에 이 함수의 "한 시트=한 날짜"
    가정으로 그대로 이어붙이면 데이터가 시트 맨 끝(전혀 다른 미래 날짜 배너
    밑)에 엉뚱하게 붙는다. 그 시트에 날짜 배너(G열, _DATE_RE 매칭)가 2개
    이상이면 옛 구조로 판단해 건드리지 않고 "{sheet_name} (마감)"이라는
    새 이름으로 대신 만든다."""
    with ExcelWriteLock():
        wb = get_workbook()
        for iso, rows in dates_payload.items():
            d = datetime.strptime(iso, "%Y-%m-%d")
            sheet_name = f"{d.month}-{d.day}"
            if sheet_name in [s.name for s in wb.sheets]:
                ws = wb.sheets[sheet_name]
                used = _com_retry(lambda: ws.used_range)
                banner_count = 0
                for r in range(1, used.last_cell.row + 1):
                    v = ws.cells(r, 7).value
                    if v and isinstance(v, str) and _DATE_RE.search(v):
                        banner_count += 1
                if banner_count > 1:
                    sheet_name = f"{sheet_name} (마감)"
            write_finalized_sheet_com(wb, sheet_name, d, rows)
        wb.save()
