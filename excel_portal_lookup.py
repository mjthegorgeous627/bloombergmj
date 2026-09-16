"""Read portal automation defaults from the live delivery Excel workbook."""

import re
from collections import Counter
from datetime import datetime

import openpyxl
import xlwings as xw

# 2026-09-01: 공유파일(EXCEL_PATH)은 이제 매일 SAP 수집이 직접 안 쓰고(사용자
# 결정 - workbench가 메인, 공유파일은 마감된 확정 데이터만 담는 아카이브),
# "지금 살아있는" 상태를 담는 파일은 board_sync.py가 20분마다 자동 채우는
# REALTIME_SYNC_PATH로 바뀌었다. 이 모듈은 POD remarks 기본값을 실시간 최신
# 상태에서 뽑아와야 하므로 그쪽을 본다 - 아래 코드는 전부 EXCEL_PATH라는
# 이름 그대로 두고 값만 바꿔서, 나머지 로직(시트 탐색 등)은 안 건드렸다.
from config import REALTIME_SYNC_PATH as EXCEL_PATH


ORDER_RE = re.compile(r"\b(\d{7,})\b")
OBD_RE = re.compile(r"\bOBD\s*[:#-]?\s*(\d{7,})\b", re.IGNORECASE)
QTY_RE = re.compile(r"\bQty\s*[:#-]?\s*(\d+)\b|\bx\s*(\d+)\b", re.IGNORECASE)


def _merged_map(ws):
    mapping = {}
    if not hasattr(ws, "merged_cells"):
        return mapping
    for merged in ws.merged_cells.ranges:
        value = ws.cell(merged.min_row, merged.min_col).value
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                mapping[(row, col)] = value
    return mapping


def _cell_value(ws, row, col, merged=None):
    if merged and (row, col) in merged:
        return merged[(row, col)]
    return ws.cell(row, col).value


def _norm_digits(value):
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def _today_sheet_names(sheetnames):
    today = datetime.today()
    names = []
    candidates = [f"{today.month}-{today.day}", f"{today.month}.{today.day}", f"{today.month}_{today.day}"]
    for name in candidates:
        if name in sheetnames and name not in names:
            names.append(name)
    return names


def _recent_sheet_names(wb, days=45):
    return _today_sheet_names(wb.sheetnames)


def _recent_sheet_names_from_list(sheetnames, days=45):
    return _today_sheet_names(sheetnames)


def _row_order_text(ws, row, merged=None):
    return str(_cell_value(ws, row, 1, merged) or "")


def _is_order_row(ws, row, merged=None):
    return bool(ORDER_RE.search(_row_order_text(ws, row, merged)))


def _find_order_row(ws, order_num, merged=None):
    target = _norm_digits(order_num).lstrip("0")
    for row in range(1, ws.max_row + 1):
        text = _row_order_text(ws, row, merged)
        for found in ORDER_RE.findall(text):
            if found.lstrip("0") == target:
                return row
    return None


def _group_rows(ws, first_row, order_num, merged=None):
    max_row = ws.max_row
    rows = [first_row]
    target = _norm_digits(order_num).lstrip("0")

    row = first_row + 1
    while row <= max_row:
        text = _row_order_text(ws, row, merged).strip()
        if text and _is_order_row(ws, row, merged):
            if target in _norm_digits(text).lstrip("0"):
                rows.append(row)
                row += 1
                continue
            break
        if text:
            break
        # Empty A cell often means continuation row for the same merged customer/address.
        if any(_cell_value(ws, row, c, merged) for c in range(2, 5)):
            rows.append(row)
            row += 1
            continue
        break
    return rows


def _item_label(description):
    text = str(description or "").upper()
    if "MONITOR" in text:
        return "Monitor"
    if "AUTHENTICATION" in text or "BUNIT" in text:
        return "BUNIT"
    if "LAPTOP" in text or "PC" in text:
        return "PC"
    if "KEYBOARD" in text or "KB" in text:
        return "KB"
    if "AUTHENTICATION" in text:
        return "BUNIT"
    return "Item"


def _extract_qty(description):
    match = QTY_RE.search(str(description or ""))
    if not match:
        return 1
    try:
        value = match.group(1) or match.group(2)
        return max(1, int(value))
    except ValueError:
        return 1


def _format_items(labels):
    counts = Counter(labels)
    return ", ".join(f"{count}x{label}" for label, count in counts.items())


def _format_item_counts(pairs):
    counts = Counter()
    for label, qty in pairs:
        try:
            qty_int = max(1, int(qty))
        except (ValueError, TypeError):
            qty_int = 1
        counts[label] += qty_int
    return ", ".join(f"{count}x{label}" for label, count in counts.items())


def _line_kind(order_text):
    text = str(order_text or "").upper()
    if "회수" in text:
        return "pickup"
    if "배송" in text:
        return "delivery"
    return "delivery"


def _order_types(rows):
    joined = "\n".join(str(row.get("order_text") or "") for row in rows).upper()
    return sorted(set(re.findall(r"\b(ZOR|SDSK|ZRX|ZRE|ZINP|ZINT|ORD)\b", joined)))


def _extract_obd(order_text):
    match = OBD_RE.search(str(order_text or ""))
    return match.group(1) if match else ""


def _filter_rows_by_delivery(rows, delivery):
    delivery_digits = _norm_digits(delivery)
    if not delivery_digits:
        return rows

    selected = []
    current_obd = ""
    for row in rows:
        row_obd = _norm_digits(row.get("obd"))
        if row_obd:
            current_obd = row_obd
        if current_obd == delivery_digits:
            selected.append(row)
    return selected


def _build_result(sheet_name, order_num, rows, pickup_mode=None, source=None):
    customer = next((r["customer"] for r in rows if r["customer"]), "")
    obd = next((r["obd"] for r in rows if r["obd"]), "")
    result = {
        "sheet": sheet_name,
        "start_row": rows[0]["row"],
        "end_row": rows[-1]["row"],
        "order": str(order_num),
        "obd": obd,
        "customer": customer,
        "remarks": build_default_remarks(rows, customer, pickup_mode=pickup_mode),
        "rows": rows,
    }
    if source:
        result["source"] = source
    return result


def build_default_remarks(rows, signed_by, pickup_mode=None):
    deliveries = []
    pickups = []
    pickup_expected = False
    for row in rows:
        label = _item_label(row.get("description"))
        qty = row.get("qty") or _extract_qty(row.get("description"))
        if _line_kind(row.get("order_text")) == "pickup":
            pickup_expected = True
            if pickup_mode == "collected":
                pickups.append((label, qty))
        else:
            deliveries.append((label, qty))

    delivered = _format_item_counts(deliveries) or "1xItem"
    collected = _format_item_counts(pickups)
    signed_by = signed_by or ""
    order_types = _order_types(rows)

    if "ZRX" in order_types or pickup_expected:
        if collected:
            return f"Delivered {delivered}, and collected {collected}, to/from {signed_by}."
        if pickup_mode == "later":
            return f"Delivered {delivered} to {signed_by}, and will collect later."
        return f"Delivered {delivered} to {signed_by}, and will collect later."
    return f"Delivered {delivered} to {signed_by}."


def _open_live_workbook():
    filename = EXCEL_PATH.split("\\")[-1]
    try:
        for book in xw.books:
            if book.name == filename:
                return book
    except Exception:
        return None
    return None


def _row_value(values, row_idx, col_idx):
    try:
        return values[row_idx][col_idx]
    except Exception:
        return None


def _find_excel_order_live(order_num, pickup_mode=None, delivery=None):
    """Read from the currently open Excel workbook, including unsaved edits."""
    wb = _open_live_workbook()
    if wb is None:
        return None

    target = _norm_digits(order_num).lstrip("0")
    sheet_names = [sheet.name for sheet in wb.sheets]
    for sheet_name in _recent_sheet_names_from_list(sheet_names):
        ws = wb.sheets[sheet_name]
        used = ws.used_range
        values = used.value
        if values is None:
            continue
        if not isinstance(values, list):
            continue
        if values and not isinstance(values[0], list):
            values = [values]

        first_idx = None
        for idx, row in enumerate(values):
            a_value = _row_value(values, idx, 0)
            if a_value and target in _norm_digits(a_value).lstrip("0"):
                first_idx = idx
                break
        if first_idx is None:
            continue

        group = [first_idx]
        idx = first_idx + 1
        while idx < len(values):
            a_text = str(_row_value(values, idx, 0) or "").strip()
            if a_text and ORDER_RE.search(a_text):
                if target in _norm_digits(a_text).lstrip("0"):
                    group.append(idx)
                    idx += 1
                    continue
                break
            if a_text:
                break
            if any(_row_value(values, idx, c) for c in range(1, 4)):
                group.append(idx)
                idx += 1
                continue
            break

        rows = []
        for idx in group:
            order_text = _row_value(values, idx, 0) or ""
            description = _row_value(values, idx, 1) or ""
            rows.append(
                {
                    "row": used.row + idx,
                    "order_text": order_text,
                    "obd": _extract_obd(order_text),
                    "description": description,
                    "qty": _extract_qty(description),
                    "material": str(_row_value(values, idx, 2) or "").strip(),
                    "serial": str(_row_value(values, idx, 3) or "").strip(),
                    "customer": str(_row_value(values, idx, 4) or "").strip(),
                    "phone": str(_row_value(values, idx, 5) or "").strip(),
                    "address": str(_row_value(values, idx, 6) or "").strip(),
                    "memo": str(_row_value(values, idx, 7) or "").strip(),
                }
            )

        rows = _filter_rows_by_delivery(rows, delivery)
        if not rows:
            continue
        return _build_result(ws.name, order_num, rows, pickup_mode=pickup_mode, source="live_excel")
    return None


def find_excel_order(order_num, pickup_mode=None, delivery=None, use_live=False):
    """Return Excel-derived portal defaults for an order number."""
    if use_live:
        try:
            live = _find_excel_order_live(order_num, pickup_mode=pickup_mode, delivery=delivery)
        except Exception:
            live = None
        if live:
            return live

    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    try:
        for sheet_name in _recent_sheet_names(wb):
            ws = wb[sheet_name]
            # Quick scan A column first; build merged map only if this sheet may match.
            target = _norm_digits(order_num).lstrip("0")
            possible = False
            for row in range(1, ws.max_row + 1):
                value = ws.cell(row, 1).value
                if value and target in _norm_digits(value).lstrip("0"):
                    possible = True
                    break
            if not possible:
                continue
            merged = _merged_map(ws)
            first_row = _find_order_row(ws, order_num, merged)
            if not first_row:
                continue
            group = _group_rows(ws, first_row, order_num, merged)
            rows = []
            for row in group:
                order_text = _cell_value(ws, row, 1, merged) or ""
                description = _cell_value(ws, row, 2, merged) or ""
                rows.append(
                    {
                        "row": row,
                        "order_text": order_text,
                        "obd": _extract_obd(order_text),
                        "description": description,
                        "qty": _extract_qty(description),
                        "material": str(_cell_value(ws, row, 3, merged) or "").strip(),
                        "serial": str(_cell_value(ws, row, 4, merged) or "").strip(),
                        "customer": str(_cell_value(ws, row, 5, merged) or "").strip(),
                        "phone": str(_cell_value(ws, row, 6, merged) or "").strip(),
                        "address": str(_cell_value(ws, row, 7, merged) or "").strip(),
                        "memo": str(_cell_value(ws, row, 8, merged) or "").strip(),
                    }
                )

            rows = _filter_rows_by_delivery(rows, delivery)
            if not rows:
                continue
            return _build_result(ws.title, order_num, rows, pickup_mode=pickup_mode)
    finally:
        wb.close()
    raise LookupError(f"Excel 배송장에서 오더 {order_num}을 찾지 못했습니다.")


if __name__ == "__main__":
    import argparse
    import json

    parser = argparse.ArgumentParser()
    parser.add_argument("order")
    parser.add_argument("--delivery")
    args = parser.parse_args()
    print(json.dumps(find_excel_order(args.order, delivery=args.delivery), ensure_ascii=False, indent=2))
